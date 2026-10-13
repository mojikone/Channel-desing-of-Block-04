"""Build template/Block04_template.docx from the SUBMITTED Block 6 report "Channels Feasibility & Hydrological Report" (R5, his
comment of 2026-10-12: title and logos as issued; the earlier template came from the working Block 6 channel design report):
the cover page (texts re-pointed to Block 04, Rev 00, the build month), the header and footer of the body sections, the
styles (Aptos body, blue headings, Caption, toc, table of figures). Everything after the cover is removed.
usage: python make_template.py [--sw]      --sw: the stormwater channels report (cover title "Hydraulic Design of Stormwater Channels", template/Block04_template_sw.docx)
"""
import sys
SW = "--sw" in sys.argv
import os, copy
from docx import Document
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = "W:/2224 Flood Protection Scheme in North & South Batinah_MAFWWR/04-DOCUMENTS/06-Design Reports/04-Feasibility Report/04-Channel Feasibility and Hydrology Report/01-B6/Channels FS and Hydrology-B6.docx"   # R5: the SUBMITTED Block 6 report (title and logos as issued)
OUT = os.path.join(HERE, "template", "Block04_template_sw.docx" if SW else "Block04_template.docx"); os.makedirs(os.path.dirname(OUT), exist_ok=True)
REPL = {"Block No. 6": "Block No. 4", "Block 6": "Block 04", "Block No 6": "Block No. 4", "April 2026": "October 2026", "May 2026": "October 2026", "Rev 00": "Rev 00"}
if SW: REPL["Channels Feasibility & Hydrological Report"] = "Hydraulic Design of Stormwater Channels"

def fix_text(root):
    for t in root.iter(qn("w:t")):
        s = t.text or ""
        for a, b in REPL.items():
            if a in s: s = s.replace(a, b)
        t.text = s

d = Document(SRC); body = d.element.body
# the body header (project title) and footer (page number, the firm's graphic) of the submitted report live in its second section:
# their references go onto the trailing section properties, which become the body section of every report built on this template
import copy as _copy
_body_sp = d.sections[1]._sectPr if len(d.sections) > 1 else None
_hf_refs = [_copy.deepcopy(c) for c in _body_sp if c.tag in (qn("w:headerReference"), qn("w:footerReference"))] if _body_sp is not None else []
kids = [k for k in body if k.tag != qn("w:sectPr")]
# the cover is everything up to and including the first paragraph that carries a section break
end = None
for i, k in enumerate(kids):
    if k.tag == qn("w:p") and k.find(qn("w:pPr")) is not None and k.find(qn("w:pPr")).find(qn("w:sectPr")) is not None:
        end = i; break
assert end is not None, "no section break after the cover"
for k in kids[end + 1:]: body.remove(k)
# drop leading empties: the two blank paragraphs above the emblem pushed the revision lines onto a second page
for k in kids[:2]:
    if k.tag == qn('w:p') and not ''.join(t.text or '' for t in k.iter(qn('w:t'))).strip(): body.remove(k)
# the trailing section: the submitted body header and footer; the cover section: none
_trail = body.find(qn("w:sectPr"))
for c in list(_trail):
    if c.tag in (qn("w:headerReference"), qn("w:footerReference")): _trail.remove(c)
for c in reversed(_hf_refs): _trail.insert(0, c)
_cover_sp = kids[end].find(qn("w:pPr")).find(qn("w:sectPr"))
for c in list(_cover_sp):
    if c.tag in (qn("w:headerReference"), qn("w:footerReference")): _cover_sp.remove(c)
print("body header/footer references:", [(c.tag.split('}')[1], c.get(qn('w:type'))) for c in _hf_refs])
fix_text(body)
if SW:   # the cover title may be split over several runs: join the runs of a paragraph that still carries the old title
    for p in body.iter(qn("w:p")):
        ts = list(p.iter(qn("w:t"))); full = "".join(t.text or "" for t in ts)
        if "Channels Feasibility" in full and len(ts) > 1:
            ts[0].text = full.replace("Channels Feasibility & Hydrological Report", "Hydraulic Design of Stormwater Channels")
            for t in ts[1:]: t.text = ""
            print("cover title joined and replaced")
for rel in list(d.part.rels.values()):   # the existing header and footer parts, without creating empty ones
    if rel.reltype.endswith("/header") or rel.reltype.endswith("/footer"): fix_text(rel.target_part.element)
# the final section properties (body pages: Letter portrait, 2.5 cm margins, the running header) stay as the trailing sectPr
# heading numbers in the heading look (comment 4): the numbering levels linked to Heading 1-3 take the heading colour, bold, size
from docx.oxml import parse_xml
num = d.part.numbering_part.element
styles = {s_.get(qn('w:styleId')): s_ for s_ in d.styles.element.findall(qn('w:style'))}
for sid, size in (('Heading1', 28), ('Heading2', 24), ('Heading3', 24)):
    st = styles.get(sid)
    if st is None: continue
    ppr = st.find(qn('w:pPr')); npr = ppr.find(qn('w:numPr')) if ppr is not None else None
    if npr is None: continue
    num_id = npr.find(qn('w:numId')).get(qn('w:val')); ilvl = (npr.find(qn('w:ilvl')).get(qn('w:val')) if npr.find(qn('w:ilvl')) is not None else '0')
    n_el = [x for x in num.findall(qn('w:num')) if x.get(qn('w:numId')) == num_id]
    if not n_el: continue
    abs_id = n_el[0].find(qn('w:abstractNumId')).get(qn('w:val'))
    for an in num.findall(qn('w:abstractNum')):
        if an.get(qn('w:abstractNumId')) != abs_id: continue
        for lvl in an.findall(qn('w:lvl')):
            if lvl.get(qn('w:ilvl')) == ilvl:
                for old in lvl.findall(qn('w:rPr')): lvl.remove(old)
                lvl.append(parse_xml(f'<w:rPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:rFonts w:asciiTheme="minorHAnsi" w:hAnsiTheme="minorHAnsi"/><w:b/><w:color w:val="0000CC"/><w:sz w:val="{size}"/><w:szCs w:val="{size}"/></w:rPr>'))
                print('numbering level styled:', sid, 'num', num_id, 'lvl', ilvl)
d.save(OUT)
d2 = Document(OUT); print("template written:", OUT, "| sections:", len(d2.sections), "| cover paragraphs:", [q.text[:40] for q in d2.paragraphs if q.text.strip()][:12])
_b = d2.sections[-1]; print("body section:", f"{_b.page_width.cm:.1f}x{_b.page_height.cm:.1f} cm", "| header linked:", _b.header.is_linked_to_previous, "| footer linked:", _b.footer.is_linked_to_previous,
      "| header:", [q.text[:70] for q in _b.header.paragraphs if q.text.strip()] if not _b.header.is_linked_to_previous else None, "| footer has PAGE:", ("PAGE" in _b.footer._element.xml) if not _b.footer.is_linked_to_previous else None)
