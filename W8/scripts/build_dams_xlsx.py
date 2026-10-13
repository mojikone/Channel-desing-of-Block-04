"""Builds W8/output/Dams_Blocks_2_4_6.xlsx from data/dams_blocks_2_4_6.csv: one sheet per block, the Block 6 table layout, tables only."""
import csv, os, openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
rows = list(csv.DictReader(open(os.path.join(ROOT, "data", "dams_blocks_2_4_6.csv"), encoding="utf-8")))
wb = openpyxl.Workbook(); wb.remove(wb.active)
hdr = ["Dam Name", "Dam Type", "Height above wadi (m)", "Crest Length (m)", "FSL (m a.s.l.)", "Reservoir Capacity (MCM)"]
thin = Side(style="thin", color="808080"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
fill = PatternFill("solid", fgColor="D9D9D9"); F = "Arial"
foot = {"2": [], "4": ["** Upgradation of the existing dam"], "6": ["* Rehabilitation of the existing dam", "** Upgradation of the existing dam"]}
for blk in ["2", "4", "6"]:
    ws = wb.create_sheet(f"Block {blk}")
    ws["A1"] = f"Characteristics of the proposed dams in Block {blk}"; ws["A1"].font = Font(name=F, size=11, bold=True)
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=3, column=j, value=h); c.font = Font(name=F, size=10, bold=True); c.fill = fill; c.border = box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 42
    r = 4
    for d in [x for x in rows if x["block"] == blk]:
        vals = [d["label"], d["type"], int(d["height_m"]), int(d["crest_length_m"]), int(d["fsl_m"]), float(d["capacity_mcm"])]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v); c.font = Font(name=F, size=10, bold=(j == 1)); c.border = box
            c.alignment = Alignment(horizontal="left" if j == 1 else "center", vertical="center")
            if j == 4: c.number_format = "#,##0"
            if j == 6: c.number_format = "0.0"
        r += 1
    for k, t in enumerate(foot[blk]):
        c = ws.cell(row=r + k, column=1, value=t); c.font = Font(name=F, size=9, italic=True)
    for j, w in enumerate([20, 14, 14, 14, 14, 16], 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.sheet_view.showGridLines = False
out = os.path.join(ROOT, "output", "Dams_Blocks_2_4_6.xlsx"); wb.save(out); print("wrote", out)
