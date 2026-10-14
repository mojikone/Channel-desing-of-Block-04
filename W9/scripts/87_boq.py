"""Step 87 (W9, engineer 2026-10-14) - one workbook with the bill of quantities and the cost estimate of the channel scheme:
channels (earthworks, protection, clearing, land of the corridors at the master plan expropriation rates), dykes, culverts, and the
summary. Every cost is a formula on a quantity and a rate; the rates sit on their own sheet so the estimate recalculates when a
rate changes. Quantities from the W9 terrain run (channels), W7 step 72 (dykes) and W9 step 71 (culverts).
    python 87_boq.py   -> W9/tables/Block 04 - BOQ and cost estimate.xlsx"""
import os, sys, ast, math
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as L
HERE = os.path.dirname(os.path.abspath(__file__)); W9 = os.path.dirname(HERE); CB = os.path.dirname(W9); W7 = os.path.join(CB, "W7")
sys.path.insert(0, os.path.join(W9, "report")); import data_facts as A   # rates and the land rule of the reports
OUT = os.path.join(W9, "tables"); os.makedirs(OUT, exist_ok=True)
HDR = PatternFill("solid", fgColor="1F5F8B"); HFONT = Font(bold=True, color="FFFFFF", name="Arial", size=10); FONT = Font(name="Arial", size=10); BOLD = Font(name="Arial", size=10, bold=True)
BLUE = Font(name="Arial", size=10, color="0000FF"); YEL = PatternFill("solid", fgColor="FFFF00")
def header(ws, labels, widths=None, height=40):
    ws.append(labels)
    for c in ws[ws.max_row]: c.fill = HDR; c.font = HFONT; c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[ws.max_row].height = height
    if widths:
        for j, w in enumerate(widths, 1): ws.column_dimensions[L(j)].width = w
wb = Workbook()
# ---- Rates (inputs in blue on yellow)
R = wb.active; R.title = "Rates"
rates = [("clear", "Clearing and grubbing", "m²", 0.30), ("exc", "Excavation", "m³", 2.10), ("fill", "Embankment and dyke fill", "m³", 4.00), ("riprap", "Riprap 0.3 m over geotextile", "m³", 13.20),
         ("conc", "Reinforced concrete of the culverts", "m³", 106.90), ("rebar", "Reinforcement (120 kg/m³)", "t", 417.10), ("blind", "Blinding concrete", "m³", 45.00), ("wp", "Waterproofing", "m²", 2.00),
         ("abc", "Aggregate base course", "m³", 6.50), ("bbc", "Bituminous base course", "m³", 45.00), ("wbc", "Wearing course", "m³", 47.00), ("pc", "Prime coat", "m²", 0.45), ("tc", "Tack coat", "m²", 0.35),
         ("mark", "Road marking", "m²", 4.00), ("rail", "Guardrail", "m", 53.00), ("post", "Guardrail post", "no.", 3.00), ("cng", "Compaction of natural ground", "m²", 0.20), ("removal", "Removal of the existing structure", "m³", 4.00)]
header(R, ["Code", "Item", "Unit", "Rate (OMR)"], [10, 44, 8, 12]); RATE_ROW = {}
for k, name, unit, v in rates:
    R.append([k, name, unit, v]); RATE_ROW[k] = R.max_row; c = R.cell(row=R.max_row, column=4); c.font = BLUE; c.fill = YEL; c.number_format = "#,##0.00"
R.append([]); R.append(["Land", "Expropriation rate by land use of the cadastre (master plan)", "", ""]); R[R.max_row][1].font = BOLD
LAND_ROW = {}
for cls, v in sorted(A.LAND_RATE_OMR_HA.items()):
    R.append([cls, cls + " land", "OMR/ha", v]); LAND_ROW[cls] = R.max_row; c = R.cell(row=R.max_row, column=4); c.font = BLUE; c.fill = YEL; c.number_format = "#,##0"
R.append(["", "Government, planned and vacant land", "OMR/ha", 0]); c = R.cell(row=R.max_row, column=4); c.font = BLUE; c.fill = YEL
R.append([]); R.append(["", "Blue on yellow: the unit rates (master plan estimate); change them here and every cost follows."]); R[R.max_row][1].font = Font(name="Arial", size=9, italic=True)
def rate(k): return f"Rates!$D${RATE_ROW[k]}"
# ---- Channels
q = pd.read_csv(os.path.join(W9, "volumes", "channel_quantities.csv")); q["land_take_ha"] = q.land_take_ha.map(lambda s: ast.literal_eval(s) if isinstance(s, str) and s.startswith("{") else {})
classes = [c for c in sorted(A.LAND_RATE_OMR_HA)]; others = sorted({A.LAND_CLASS.get(str(k).strip(), "") for lt in q.land_take_ha for k in lt} - set(classes) - {""})
C = wb.create_sheet("Channels")
cols = ["Channel", "Class", "Length (m)", "Clearing (m²)", "Excavation (m³)", "Fill (m³)", "Riprap area (m²)", "Riprap volume 0.3 m (m³)", "Clearing cost (OMR)", "Excavation cost (OMR)", "Fill cost (OMR)", "Riprap cost (OMR)", "Construction cost (OMR)", "Corridor (ha)"] + [f"{c} land (ha)" for c in classes] + ["Other land, no cost (ha)", "Land cost (OMR)", "Channel cost (OMR)"]
header(C, cols, [13, 7, 10, 12, 12, 12, 12, 12, 13, 13, 13, 13, 14, 10] + [11] * len(classes) + [11, 13, 14])
CL = {name: L(i + 1) for i, name in enumerate(cols)}; first = 2
for _, r in q.iterrows():
    row = C.max_row + 1; lt = {A.LAND_CLASS.get(str(k).strip(), ""): 0 for k in r.land_take_ha}
    for k, v in r.land_take_ha.items(): lt[A.LAND_CLASS.get(str(k).strip(), "")] = lt.get(A.LAND_CLASS.get(str(k).strip(), ""), 0) + float(v)
    other = sum(v for k, v in lt.items() if k not in classes)
    C.append([r.channel, r.cls, round(r.length_m, 1), round(r.footprint_m2), round(r.cut_m3), round(r.fill_m3), round(r.protection_m2), f"={CL['Riprap area (m²)']}{row}*0.3",
              f"={CL['Clearing (m²)']}{row}*{rate('clear')}", f"={CL['Excavation (m³)']}{row}*{rate('exc')}", f"={CL['Fill (m³)']}{row}*{rate('fill')}", f"={CL['Riprap volume 0.3 m (m³)']}{row}*{rate('riprap')}",
              f"=SUM({CL['Clearing cost (OMR)']}{row}:{CL['Riprap cost (OMR)']}{row})", round(r.corridor_ha, 2)] + [round(lt.get(c, 0.0), 2) for c in classes] + [round(other, 2),
              "=" + "+".join(f"{CL[f'{c} land (ha)']}{row}*Rates!$D${LAND_ROW[c]}" for c in classes), f"={CL['Construction cost (OMR)']}{row}+{CL['Land cost (OMR)']}{row}"])
last = C.max_row; C.append(["Total", ""] + [f"=SUM({L(j)}{first}:{L(j)}{last})" for j in range(3, len(cols) + 1)])
for c in C[C.max_row]: c.font = BOLD
for row in C.iter_rows(min_row=2):
    for c in row:
        if c.font != BOLD: c.font = FONT
        if c.column >= 3: c.number_format = "#,##0" if "ha" not in cols[c.column - 1] else "#,##0.00"
C.freeze_panes = "C2"
# ---- Dykes
d = pd.read_csv(os.path.join(W7, "dykes", "dykes.csv"))
D = wb.create_sheet("Dykes")
dcols = ["Dyke", "Class", "Length (m)", "Height (m)", "Clearing (m²)", "Excavation (m³)", "Fill (m³)", "Riprap volume (m³)", "Clearing cost (OMR)", "Excavation cost (OMR)", "Fill cost (OMR)", "Riprap cost (OMR)", "Dyke cost (OMR)"]
header(D, dcols, [13, 10, 10, 9, 12, 12, 12, 12, 13, 13, 13, 13, 14]); DL = {name: L(i + 1) for i, name in enumerate(dcols)}
for _, r in d.iterrows():
    row = D.max_row + 1
    D.append([r["name"], r["cls"], round(r.L_m, 1), round(r.H_m, 1), round(r.clearing_m2), round(r.excavation_m3), round(r.fill_m3), round(r.riprap_m3),
              f"={DL['Clearing (m²)']}{row}*{rate('clear')}", f"={DL['Excavation (m³)']}{row}*{rate('exc')}", f"={DL['Fill (m³)']}{row}*{rate('fill')}", f"={DL['Riprap volume (m³)']}{row}*{rate('riprap')}", f"=SUM({DL['Clearing cost (OMR)']}{row}:{DL['Riprap cost (OMR)']}{row})"])
dl = D.max_row; D.append(["Total", ""] + [f"=SUM({L(j)}2:{L(j)}{dl})" for j in range(3, len(dcols) + 1)])
for c in D[D.max_row]: c.font = BOLD
for row in D.iter_rows(min_row=2):
    for c in row:
        if c.font != BOLD: c.font = FONT
        if c.column >= 3: c.number_format = "#,##0"
D.freeze_panes = "C2"
# ---- Culverts: the item quantities of the W9 build-up, the costs by formula
p = pd.read_csv(os.path.join(W9, "culverts", "culverts_proposed.csv"), keep_default_na=False)
WALL = 0.25
K = wb.create_sheet("Culverts")
kcols = ["ID", "Road", "Channel", "Reach", "Cells", "Cell width (m)", "Cell height (m)", "Overall width G (m)", "Length L (m)", "Concrete (m³)", "Reinforcement (t)", "Blinding (m³)", "Waterproofing (m²)", "Pavement area (m²)", "Aggregate base (m³)", "Bituminous base (m³)", "Wearing course (m³)", "Prime coat (m²)", "Tack coat (m²)", "Marking (m²)", "Guardrail (m)", "Posts (no.)", "Compaction (m²)", "Removal of existing (m³)",
         "Concrete cost", "Reinforcement cost", "Blinding cost", "Waterproofing cost", "Aggregate base cost", "Bituminous base cost", "Wearing course cost", "Prime coat cost", "Tack coat cost", "Marking cost", "Guardrail cost", "Posts cost", "Compaction cost", "Removal cost", "Culvert cost (OMR)"]
header(K, kcols, [5, 16, 12, 12, 6, 8, 8, 10, 8] + [11] * 15 + [12] * 15, 52); KL = {name: L(i + 1) for i, name in enumerate(kcols)}
for _, r in p.iterrows():
    N = int(r.cells); B = float(r.cell_width_m); Dh = float(r.cell_height_m); Lc = float(r.length_m); G = N * B + (N + 1) * WALL; existing = str(r.existing_type) not in ("None", "", "nan")
    conc = (N * ((B + WALL) * (Dh + 0.6) - B * Dh) + 0.8) * Lc; PC = G * Lc * 1.1; row = K.max_row + 1
    K.append([int(r.id), r.road, r.channel, r.reach, N, B, Dh, round(G, 2), Lc, round(conc, 1), f"=0.12*{KL['Concrete (m³)']}{row}", round(G * Lc * 0.1 * 1.1, 1), round(2 * Dh * Lc + B * Lc, 1), round(PC, 1), round(PC * 0.1, 1), round(PC * 0.05, 1), round(PC * 0.05, 1), round(PC, 1), round(2 * PC, 1), round(G * 0.1 * 1.1, 2), round(2 * G * 1.1, 1), math.ceil(2 * G * 1.1 / 4), round((Lc + 3) * G * 1.1, 1), round(G * Lc * (Dh + 0.5) * 1.1, 1) if existing else 0,
              f"={KL['Concrete (m³)']}{row}*{rate('conc')}", f"={KL['Reinforcement (t)']}{row}*{rate('rebar')}", f"={KL['Blinding (m³)']}{row}*{rate('blind')}", f"={KL['Waterproofing (m²)']}{row}*{rate('wp')}", f"={KL['Aggregate base (m³)']}{row}*{rate('abc')}", f"={KL['Bituminous base (m³)']}{row}*{rate('bbc')}", f"={KL['Wearing course (m³)']}{row}*{rate('wbc')}",
              f"={KL['Prime coat (m²)']}{row}*{rate('pc')}", f"={KL['Tack coat (m²)']}{row}*{rate('tc')}", f"={KL['Marking (m²)']}{row}*{rate('mark')}", f"={KL['Guardrail (m)']}{row}*{rate('rail')}", f"={KL['Posts (no.)']}{row}*{rate('post')}", f"={KL['Compaction (m²)']}{row}*{rate('cng')}", f"={KL['Removal of existing (m³)']}{row}*{rate('removal')}",
              f"=SUM({KL['Concrete cost']}{row}:{KL['Removal cost']}{row})"])
kl = K.max_row; K.append(["Total", "", "", ""] + [f"=SUM({L(j)}2:{L(j)}{kl})" for j in range(5, len(kcols) + 1)])
for c in K[K.max_row]: c.font = BOLD
for row in K.iter_rows(min_row=2):
    for c in row:
        if c.font != BOLD: c.font = FONT
        if c.column >= 5: c.number_format = "#,##0"
K.freeze_panes = "E2"
# ---- Summary
S = wb.create_sheet("Summary", 0)
S["A1"] = "Block 04 - bill of quantities and cost estimate of the channel scheme (200-year design)"; S["A1"].font = Font(name="Arial", size=12, bold=True)
header(S, ["Item", "Quantity", "Unit", "Cost (OMR)", "Cost (million OMR)"], [46, 14, 8, 16, 14]); S.row_dimensions[2].height = 20
tot = C.max_row; dtot = D.max_row; ktot = K.max_row
lines = [("Channels - clearing and grubbing", f"=Channels!{CL['Clearing (m²)']}{tot}", "m²", f"=Channels!{CL['Clearing cost (OMR)']}{tot}"), ("Channels - excavation", f"=Channels!{CL['Excavation (m³)']}{tot}", "m³", f"=Channels!{CL['Excavation cost (OMR)']}{tot}"),
         ("Channels - fill", f"=Channels!{CL['Fill (m³)']}{tot}", "m³", f"=Channels!{CL['Fill cost (OMR)']}{tot}"), ("Channels - riprap", f"=Channels!{CL['Riprap volume 0.3 m (m³)']}{tot}", "m³", f"=Channels!{CL['Riprap cost (OMR)']}{tot}"),
         ("Channels - construction", f"=Channels!{CL['Length (m)']}{tot}", "m", f"=Channels!{CL['Construction cost (OMR)']}{tot}"), ("Channels - land of the corridors (expropriation)", f"=Channels!{CL['Corridor (ha)']}{tot}", "ha", f"=Channels!{CL['Land cost (OMR)']}{tot}"),
         ("Dykes", f"=Dykes!{DL['Length (m)']}{dtot}", "m", f"=Dykes!{DL['Dyke cost (OMR)']}{dtot}"), ("Culverts", f"=Culverts!{KL['Cells']}{ktot}", "cells", f"=Culverts!{KL['Culvert cost (OMR)']}{ktot}")]
for name, qf, unit, cf in lines:
    S.append([name, qf, unit, cf, f"=D{S.max_row + 1}/1000000"])
S.append(["Total", "", "", f"=D{S.max_row - 2}+D{S.max_row - 1}+D{S.max_row}+D{S.max_row - 3}", f"=D{S.max_row + 1}/1000000"])
for c in S[S.max_row]: c.font = BOLD
for row in S.iter_rows(min_row=3):
    for c in row:
        if c.font != BOLD: c.font = FONT
    row[1].number_format = "#,##0"; row[3].number_format = "#,##0"; row[4].number_format = "#,##0.00"
S.append([]); S.append(["The total is the construction of the channels, the land of their corridors, the dykes and the culverts. Quantities: channels measured on the designed sections against the 0.5 m terrain model; dykes on the typical fill section at the adopted height; culverts by the concrete, reinforcement, blinding, waterproofing and road works over them. Rates: sheet Rates (master plan estimate)."])
S[S.max_row][0].font = Font(name="Arial", size=9, italic=True)
path = os.path.join(OUT, "Block 04 - BOQ and cost estimate.xlsx"); wb.save(path); print("written", path, "| channels", len(q), "dykes", len(d), "culverts", len(p))
