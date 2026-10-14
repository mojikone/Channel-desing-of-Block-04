"""Step 85 (W9, engineer 2026-10-14) - the bypass road design package, two workbooks:
  * Block 04 - Flows and hydrographs at the bypass crossings.xlsx: for every crossing of a channel with the planned North Batinah
    Bypass (W9 culverts: 17 crossings, 9 channels starting just downstream of it), the peak flow and volume of all 13 return periods
    and the 48-hour hydrographs of the HEC-HMS element that feeds the channel just upstream of the crossing (dams in place), with the
    charts of the dam hydrograph workbook (one sheet per crossing); the 200-year design flow of the reach for comparison.
  * Block 04 - Culverts at the bypass crossings.xlsx: the proposed box culverts under the bypass with the channel data at the crossing.
    python 85_bypass_flows.py"""
import os, sys, re
import numpy as np, pandas as pd, geopandas as gpd, hecdss
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.data_source import StrRef
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter as L
HERE = os.path.dirname(os.path.abspath(__file__)); W9 = os.path.dirname(HERE); CB = os.path.dirname(W9); B4 = os.path.dirname(CB); W3 = os.path.join(CB, "W3"); W6 = os.path.join(CB, "W6")
sys.path.insert(0, os.path.join(W3, "scripts")); from xlstyle import style_chart   # the chart look of the dam workbook
HMS = os.path.join(B4, "Hydrology", "HEC_HMS_Block_04"); OUT = os.path.join(W9, "bypass"); os.makedirs(OUT, exist_ok=True)
RPS = ["2yr", "5yr", "10yr", "25yr", "50yr", "100yr", "200yr", "500yr", "1000yr", "2000yr", "5000yr", "10000yr", "PMP"]
LAB = {rp: (rp.replace("yr", "-year") if rp != "PMP" else "PMP") for rp in RPS}
RAMP = ["C6DBEF", "B3CDE8", "9EC9E2", "86B9DC", "6BAED6", "4F9BCF", "3E8AC5", "3182BD", "2171B5", "1A5FA3", "0D4A8E", "08306B"]; DT_S = 300
HDR = PatternFill("solid", fgColor="1F5F8B"); HFONT = Font(bold=True, color="FFFFFF", name="Arial", size=10); FONT = Font(name="Arial", size=10)
cr = pd.read_csv(os.path.join(W9, "culverts", "crossings.csv")); by = cr[cr.road == "North Batinah Bypass"].copy()
pr = pd.read_csv(os.path.join(W9, "culverts", "culverts_proposed.csv"), keep_default_na=False); prb = pr[pr.road == "North Batinah Bypass"].copy()
pts = gpd.read_file(os.path.join(W6, "final", "Block4_design_points.shp"))
def element_for(channel, station):
    """the design point with an inflow element nearest upstream of the station (else the first one downstream)"""
    p = pts[(pts.Channel == channel) & pts.Flow.notna()].sort_values("Station")
    up = p[p.Station <= station + 1.0]
    row = up.iloc[-1] if len(up) else (p.iloc[0] if len(p) else None)
    return (None, None, None) if row is None else (str(row.Flow), float(row.Station), str(row.SGLabel))
def own_point(path):
    parts = path.split("/"); return len(parts) > 6 and ">" in parts[6] and parts[6].split(">")[-1] == parts[2]
# the DSS flows of the elements, per return period
elements = {}
for _, k in by.iterrows():
    el, sta, lab = element_for(k.channel, k.station_m); elements[k.id] = (el, sta, lab)
need = sorted({e for e, _, _ in elements.values() if e})
series, peaks, vols = {}, {}, {}
for rp in RPS:
    f = os.path.join(HMS, f"DA_{rp}_B6.dss")
    if not os.path.exists(f): print("MISSING", f); continue
    d = hecdss.HecDss(f); paths = [str(p) for p in d.get_catalog()]
    want = {p.split("/")[2]: p for p in paths if own_point(p) and p.split("/")[3] == "FLOW"}
    for el in need:
        p = want.get(el)
        if p is None: print("no FLOW path for", el, rp); continue
        v = np.asarray(d.get(p).values, dtype=float); series[(el, rp)] = v; peaks[(el, rp)] = float(v.max()); vols[(el, rp)] = float(v.sum() * DT_S / 1e6)
    d.close(); print(rp, "read")
# ---- workbook 1: flows and hydrographs
wb = Workbook(); ws = wb.active; ws.title = "Crossings"
head = ["ID", "Channel", "Class", "Reach at the crossing", "Station (m)", "Easting", "Northing", "Skew (deg)", "Design flow 200-year, reach (m³/s)", "Bed width 200-year (m)", "Depth (m)", "Slope (m/m)", "HEC-HMS element", "Element station (m)", "Design point"] + [f"Qp {LAB[rp]} (m³/s)" for rp in RPS] + [f"Volume {LAB[rp]} (Mm³)" for rp in RPS]
ws.append(head)
for c in ws[1]: c.fill = HDR; c.font = HFONT; c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
ws.row_dimensions[1].height = 48
for _, k in by.sort_values("id").iterrows():
    el, sta, lab = elements[k.id]
    ws.append([int(k.id), k.channel, k.cls, k.reach, round(k.station_m, 1), round(k.x, 1), round(k.y, 1), k.skew_deg, round(k.Q200, 1), k.span_m, k.rise_m, round(k.slope, 5), el, (round(sta, 1) if sta is not None else None), lab]
              + [round(peaks.get((el, rp), float("nan")), 1) for rp in RPS] + [round(vols.get((el, rp), float("nan")), 3) for rp in RPS])
for row in ws.iter_rows(min_row=2):
    for c in row: c.font = FONT
for j, w in enumerate([5, 13, 7, 14, 10, 10, 11, 8, 12, 10, 8, 9, 14, 11, 12] + [10] * 26, 1): ws.column_dimensions[L(j)].width = w
ws.freeze_panes = "B2"
def ser(ws_, xref, yref, title_ref, colour, width_emu=15875, dash=None):
    s = Series(yref, xref, title=None); s.tx = SeriesLabel(strRef=StrRef(title_ref)); s.marker.symbol = "none"; s.smooth = False
    s.graphicalProperties.line.solidFill = colour; s.graphicalProperties.line.width = width_emu; s.graphicalProperties.line.cap = "rnd"
    if dash: s.graphicalProperties.line.dashStyle = dash
    return s
for _, k in by.sort_values("id").iterrows():
    el, sta, lab = elements[k.id]
    if not el: continue
    name = re.sub(r"[\\/*?:\[\]]", "", f"{int(k.id):02d} {k.channel}")[:31]; w2 = wb.create_sheet(name)
    w2["A1"] = f"{k.channel} at the North Batinah Bypass, station {k.station_m:.0f} m - hydrographs of HEC-HMS element {el} (dams in place), 5-minute values"; w2["A1"].font = Font(bold=True, name="Arial", size=11)
    w2.append(["Time (h)"] + [f"Q {LAB[rp]} (m³/s)" for rp in RPS])
    for c in w2[2]: c.fill = HDR; c.font = HFONT; c.alignment = Alignment(horizontal="center", wrap_text=True)
    n = max(len(series.get((el, rp), [])) for rp in RPS); n = min(n, 577)
    for i in range(n):
        w2.append([round(i * DT_S / 3600, 4)] + [round(float(series[(el, rp)][i]), 3) if (el, rp) in series and i < len(series[(el, rp)]) else None for rp in RPS])
    first, last = 3, 2 + n
    # the legend table: return period, peak, volume (the series titles point at the first column)
    ST = len(RPS) + 3; w2.cell(row=2, column=ST, value="Return period"); w2.cell(row=2, column=ST + 1, value="Peak (m³/s)"); w2.cell(row=2, column=ST + 2, value="Volume (Mm³)")
    for c in (w2.cell(row=2, column=ST + j) for j in range(3)): c.fill = HDR; c.font = HFONT
    for j, rp in enumerate(RPS):
        w2.cell(row=3 + j, column=ST, value=f"{LAB[rp]}  {peaks.get((el, rp), 0):,.0f} m³/s"); w2.cell(row=3 + j, column=ST + 1, value=round(peaks.get((el, rp), 0), 1)); w2.cell(row=3 + j, column=ST + 2, value=round(vols.get((el, rp), 0), 3))
    ch = ScatterChart(); ch.title = f"{k.channel} at the bypass - {el}"; ch.style = 2; ch.width = 24; ch.height = 12
    ch.x_axis.title = "Time from start of storm (h)"; ch.y_axis.title = "Flow (m³/s)"; ch.x_axis.delete = False; ch.y_axis.delete = False
    ch.x_axis.scaling.min = 0; ch.x_axis.scaling.max = 48; ch.x_axis.majorUnit = 6; ch.y_axis.scaling.min = 0; ch.x_axis.number_format = "0"; ch.y_axis.number_format = "#,##0"; ch.x_axis.majorGridlines = None
    style_chart(ch); xs = Reference(w2, min_col=1, min_row=first, max_row=last)
    for j, rp in enumerate(RPS):
        if (el, rp) not in series: continue
        colour = "000000" if rp == "PMP" else RAMP[min(j, len(RAMP) - 1)]
        ch.series.append(ser(w2, xs, Reference(w2, min_col=2 + j, min_row=first, max_row=last), f"'{name}'!${L(ST)}${3 + j}", colour, 15875 if rp != "PMP" else 19050, "dash" if rp == "PMP" else None))
    ch.legend.position = "r"; w2.add_chart(ch, f"{L(ST + 4)}2")
    w2.column_dimensions["A"].width = 10
    for j in range(len(RPS)): w2.column_dimensions[L(2 + j)].width = 11
    w2.column_dimensions[L(ST)].width = 22; w2.freeze_panes = "B3"
rd = wb.create_sheet("Read me", 0)
for i, t in enumerate(["Block 04 - flows and hydrographs at the crossings of the planned North Batinah Bypass", "",
                       "Crossings: every channel of the Block 04 scheme that crosses the bypass alignment, and the nine channels that start just downstream of it (their runoff comes from across the road).",
                       "Design flow 200-year: the flow of the channel reach at the crossing in the hydraulic design (200-year, dams in place).",
                       "HEC-HMS element: the junction that feeds the channel immediately upstream of the crossing; its peak flow, runoff volume and 48-hour hydrograph are given for the 13 return periods up to the PMP (depth-area analyses, dams in place).",
                       "Hydrographs: one sheet per crossing, 5-minute values, with the chart of all return periods (the darker the line, the rarer the event; PMP black dashed).",
                       "Coordinates: WGS 84 / UTM zone 40N. Stations from the channel start."], 1):
    rd.cell(row=i + 1, column=1, value=t).font = Font(name="Arial", size=10, bold=(i == 0))
rd.column_dimensions["A"].width = 160
p1 = os.path.join(OUT, "Block 04 - Flows and hydrographs at the bypass crossings.xlsx"); wb.save(p1); print("written", p1, "| crossings", len(by), "| elements", len(need))
# ---- workbook 2: the culverts under the bypass
cols = [("id", "ID"), ("channel", "Channel"), ("reach", "Reach"), ("station_m", "Station (m)"), ("Q200", "Design flow 200-year (m³/s)"), ("channel_width_m", "Channel bed width (m)"), ("channel_depth_m", "Channel depth (m)"), ("slope", "Slope (m/m)"), ("skew_deg", "Skew (deg)"),
        ("cells", "Cells"), ("cell_width_m", "Cell width (m)"), ("cell_height_m", "Cell height (m)"), ("size", "Cells x width x height (m)"), ("overall_width_m", "Overall width (m)"), ("length_m", "Length across the road (m)"), ("HW_m", "Headwater (m)"), ("HW_over_D", "Headwater / height"), ("control", "Control"),
        ("tailwater_m", "Tailwater (m)"), ("concrete_m3", "Concrete (m³)"), ("rebar_t", "Reinforcement (t)"), ("cost_OMR", "Cost (OMR)")]
xy = by.set_index("id")[["x", "y"]]
wb2 = Workbook(); w = wb2.active; w.title = "Culverts under the bypass"
w.append([h for _, h in cols] + ["Easting", "Northing"])
for c in w[1]: c.fill = HDR; c.font = HFONT; c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
w.row_dimensions[1].height = 48
for _, r in prb.sort_values("id").iterrows():
    w.append([r[k] if str(r[k]) not in ("", "nan") else None for k, _ in cols] + [round(float(xy.loc[r.id, "x"]), 1), round(float(xy.loc[r.id, "y"]), 1)])
for row in w.iter_rows(min_row=2):
    for c in row: c.font = FONT
for j in range(1, len(cols) + 3): w.column_dimensions[L(j)].width = 12
w.freeze_panes = "C2"
rd2 = wb2.create_sheet("Read me", 0)
for i, t in enumerate(["Block 04 - proposed box culverts under the planned North Batinah Bypass", "",
                       "One box culvert per crossing: cells of 4.0 m clear width and the full channel depth as clear height (2.5 m, or 3.0 m where the channel is deeper), the number of cells from the designed bed width of the channel reach at the crossing, the barrel on the slope of the channel; length across the road = the bypass corridor taken as 45 m.",
                       "Headwater: inlet- and outlet-control check (FHWA culvert design method, square-edged inlet with wingwalls, entrance loss 0.4, barrel roughness 0.015, tailwater at the normal depth of the channel downstream); every culvert keeps its headwater below 1.2 times the clear height.",
                       "Cost: concrete and reinforcement, blinding, waterproofing, the road works over the culvert, at the master plan unit rates (OMR)."], 1):
    rd2.cell(row=i + 1, column=1, value=t).font = Font(name="Arial", size=10, bold=(i == 0))
rd2.column_dimensions["A"].width = 160
p2 = os.path.join(OUT, "Block 04 - Culverts at the bypass crossings.xlsx"); wb2.save(p2); print("written", p2, "| culverts", len(prb), "| cost %.2f mOMR" % (prb.cost_OMR.astype(float).sum() / 1e6))
