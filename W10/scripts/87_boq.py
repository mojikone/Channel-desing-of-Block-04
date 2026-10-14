"""Step 87 (W10; W9 origin) - the one BOQ and cost estimate workbook of the channel scheme (his request, 2026-10-14): the
channels (earthworks, protection, clearing, land of the corridors at the master plan expropriation rates), dykes, culverts, and the
summary. Every cost is a formula on a quantity and a rate; the rates sit on their own sheet so the estimate recalculates when a
rate changes. W10 (his instruction, 2026-10-14, "formulas live inside the cells"): the dyke quantities follow by formula from the
length and the height (the typical fill section), the culvert quantities from the cells, the width, the depth and the length; a
catchment column on every sheet and a by-catchment block on the summary.
Quantities from the W9 terrain run (channels, measured values), W7 step 72 (dykes) and W10 step 71 (culverts).
    python 87_boq.py"""
import os, sys, math, ast
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
HERE = os.path.dirname(os.path.abspath(__file__)); W9 = os.path.dirname(HERE); CB = os.path.dirname(W9); W7 = os.path.join(CB, "W7"); W9P = os.path.join(CB, "W9")
sys.path.insert(0, os.path.join(W9, "report")); import data_facts as A   # rates and the land rule of the reports
OUT = os.path.join(W9, "tables"); os.makedirs(OUT, exist_ok=True)
HDR = PatternFill("solid", fgColor="1F5F8B"); HFONT = Font(bold=True, color="FFFFFF", name="Arial", size=10); FONT = Font(name="Arial", size=10); BOLD = Font(name="Arial", size=10, bold=True)
BLUE = Font(name="Arial", size=10, color="0000FF"); YEL = PatternFill("solid", fgColor="FFFF00")
GCA = {"HAW": A.GCA_NAME["HAW"], "SHA": A.GCA_NAME["SHA"], "SAR": A.GCA_NAME["SAR"], "SAK": A.GCA_NAME["SAK"]}
def catchment(name): return GCA.get(str(name)[:3], "")
def header(ws, labels, widths=None, height=40):
    ws.append(labels)
    for c in ws[ws.max_row]: c.fill = HDR; c.font = HFONT; c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[ws.max_row].height = height
    if widths:
        for j, w in enumerate(widths, 1): ws.column_dimensions[L(j)].width = w
wb = Workbook()
# ---- Rates and constants (inputs in blue on yellow)
R = wb.active; R.title = "Rates"
rates = [("clear", "Clearing and grubbing", "m²", 0.30), ("exc", "Excavation", "m³", 2.10), ("fill", "Embankment and dyke fill", "m³", 4.00), ("riprap", "Riprap 0.3 m over geotextile", "m³", 13.20),
         ("conc", "Reinforced concrete of the culverts", "m³", 106.90), ("rebar", "Reinforcement (120 kg/m³)", "t", 417.10), ("blind", "Blinding concrete", "m³", 45.00), ("wp", "Waterproofing", "m²", 2.00),
         ("abc", "Aggregate base course", "m³", 6.50), ("bbc", "Bituminous base course", "m³", 45.00), ("wbc", "Wearing course", "m³", 47.00), ("pc", "Prime coat", "m²", 0.45), ("tc", "Tack coat", "m²", 0.35),
         ("mark", "Road marking", "m²", 4.00), ("rail", "Guardrail", "m", 53.00), ("post", "Guardrail post", "no.", 3.00), ("cng", "Compaction of natural ground", "m²", 0.20), ("removal", "Removal of the existing structure", "m³", 4.00)]
header(R, ["Code", "Item", "Unit", "Rate (OMR)"], [10, 48, 8, 12]); RATE_ROW = {}
for k, name, unit, v in rates:
    R.append([k, name, unit, v]); RATE_ROW[k] = R.max_row; c = R.cell(row=R.max_row, column=4); c.font = BLUE; c.fill = YEL; c.number_format = "#,##0.00"
R.append([]); R.append(["Land", "Expropriation rate by land use of the cadastre (master plan)", "", ""]); R[R.max_row][1].font = BOLD
LAND_ROW = {}
for cls, v in sorted(A.LAND_RATE_OMR_HA.items()):
    R.append([cls, cls + " land", "OMR/ha", v]); LAND_ROW[cls] = R.max_row; c = R.cell(row=R.max_row, column=4); c.font = BLUE; c.fill = YEL; c.number_format = "#,##0"
R.append(["", "Government, planned and vacant land", "OMR/ha", 0]); c = R.cell(row=R.max_row, column=4); c.font = BLUE; c.fill = YEL
R.append([]); R.append(["Constants", "Typical sections", "", ""]); R[R.max_row][1].font = BOLD
consts = [("riprap_t", "Riprap thickness on the channel slopes (m)", "m", 0.30), ("wall", "Culvert wall thickness (m)", "m", 0.25), ("crest", "Dyke crest width (m)", "m", 4.0), ("strip", "Dyke: stripping depth under the footprint (m)", "m", 0.4)]
CONST_ROW = {}
for k, name, unit, v in consts:
    R.append([k, name, unit, v]); CONST_ROW[k] = R.max_row; c = R.cell(row=R.max_row, column=4); c.font = BLUE; c.fill = YEL
R.append([]); R.append(["", "Blue on yellow: the unit rates (master plan estimate) and the section constants; change them here and every quantity and cost follows."]); R[R.max_row][1].font = Font(name="Arial", size=9, italic=True)
def rate(k): return f"Rates!$D${RATE_ROW[k]}"
def const(k): return f"Rates!$D${CONST_ROW[k]}"
# ---- Channels: the measured quantities (values), the riprap volume, the costs and the land cost by formula
q = pd.read_csv(os.path.join(W9P, "volumes", "channel_quantities.csv")); q["land_take_ha"] = q.land_take_ha.map(lambda s: ast.literal_eval(s) if isinstance(s, str) and s.startswith("{") else {})
classes = [c for c in sorted(A.LAND_RATE_OMR_HA)]
C = wb.create_sheet("Channels")
cols = ["Channel", "Class", "Catchment", "Length (m)", "Clearing (m²)", "Excavation (m³)", "Fill (m³)", "Riprap area (m²)", "Riprap volume (m³)", "Clearing cost (OMR)", "Excavation cost (OMR)", "Fill cost (OMR)", "Riprap cost (OMR)", "Construction cost (OMR)", "Corridor (ha)"] + [f"{c} land (ha)" for c in classes] + ["Other land, no cost (ha)", "Land cost (OMR)", "Channel cost (OMR)"]
header(C, cols, [13, 7, 14, 10, 12, 12, 12, 12, 12, 13, 13, 13, 13, 14, 10] + [11] * len(classes) + [11, 13, 14])
CL = {name: L(i + 1) for i, name in enumerate(cols)}; first = 2
for _, r in q.iterrows():
    row = C.max_row + 1; lt = {}
    for k, v in r.land_take_ha.items(): kk = A.LAND_CLASS.get(str(k).strip(), ""); lt[kk] = lt.get(kk, 0) + float(v)
    other = sum(v for k, v in lt.items() if k not in classes)
    C.append([r.channel, r.cls, catchment(r.channel), round(r.length_m, 1), round(r.footprint_m2), round(r.cut_m3), round(r.fill_m3), round(r.protection_m2), f"={CL['Riprap area (m²)']}{row}*{const('riprap_t')}",
              f"={CL['Clearing (m²)']}{row}*{rate('clear')}", f"={CL['Excavation (m³)']}{row}*{rate('exc')}", f"={CL['Fill (m³)']}{row}*{rate('fill')}", f"={CL['Riprap volume (m³)']}{row}*{rate('riprap')}",
              f"=SUM({CL['Clearing cost (OMR)']}{row}:{CL['Riprap cost (OMR)']}{row})", round(r.corridor_ha, 2)] + [round(lt.get(c, 0.0), 2) for c in classes] + [round(other, 2),
              "=" + "+".join(f"{CL[f'{c} land (ha)']}{row}*Rates!$D${LAND_ROW[c]}" for c in classes), f"={CL['Construction cost (OMR)']}{row}+{CL['Land cost (OMR)']}{row}"])
last = C.max_row; C.append(["Total", "", ""] + [f"=SUM({L(j)}{first}:{L(j)}{last})" for j in range(4, len(cols) + 1)])
for c in C[C.max_row]: c.font = BOLD
for row in C.iter_rows(min_row=2):
    for c in row:
        if c.font != BOLD: c.font = FONT
        if c.column >= 4: c.number_format = "#,##0" if "ha" not in cols[c.column - 1] else "#,##0.00"
C.freeze_panes = "D2"
# ---- Dykes: the length and the height are the inputs; the quantities of the typical fill section and the costs by formula
d = pd.read_csv(os.path.join(W7, "dykes", "dykes.csv"))
D = wb.create_sheet("Dykes")
dcols = ["Dyke", "Class", "Catchment", "Length (m)", "Height (m)", "Footprint width (m)", "Clearing (m²)", "Excavation (m³)", "Fill (m³)", "Riprap volume (m³)", "Clearing cost (OMR)", "Excavation cost (OMR)", "Fill cost (OMR)", "Riprap cost (OMR)", "Dyke cost (OMR)"]
header(D, dcols, [13, 10, 14, 10, 9, 11, 12, 12, 12, 12, 13, 13, 13, 13, 14]); DL = {name: L(i + 1) for i, name in enumerate(dcols)}
for _, r in d.iterrows():
    row = D.max_row + 1; Lm = f"{DL['Length (m)']}{row}"; H = f"{DL['Height (m)']}{row}"; F = f"{DL['Footprint width (m)']}{row}"
    D.append([r["name"], r["cls"], catchment(r["name"]), round(r.L_m, 1), round(r.H_m, 1), f"={const('crest')}+3*{H}+3", f"={F}*{Lm}", f"={const('strip')}*{F}*{Lm}", f"=({H}*({const('crest')}+1.5*{H})+{const('strip')}*{F})*{Lm}", f"=(0.525+0.2)*1.803*{H}*{Lm}",
              f"={DL['Clearing (m²)']}{row}*{rate('clear')}", f"={DL['Excavation (m³)']}{row}*{rate('exc')}", f"={DL['Fill (m³)']}{row}*{rate('fill')}", f"={DL['Riprap volume (m³)']}{row}*{rate('riprap')}", f"=SUM({DL['Clearing cost (OMR)']}{row}:{DL['Riprap cost (OMR)']}{row})"])
dl = D.max_row; D.append(["Total", "", ""] + [f"=SUM({L(j)}2:{L(j)}{dl})" for j in range(4, len(dcols) + 1)])
for c in D[D.max_row]: c.font = BOLD
for row in D.iter_rows(min_row=2):
    for c in row:
        if c.font != BOLD: c.font = FONT
        if c.column >= 4: c.number_format = "#,##0" if c.column != 5 else "0.0"
D.freeze_panes = "D2"
# ---- Culverts: cells, width, depth and length are the inputs; the overall width, the item quantities and the costs by formula
p = pd.read_csv(os.path.join(W9, "culverts", "culverts_proposed.csv"), keep_default_na=False)
K = wb.create_sheet("Culverts")
kcols = ["ID", "Road", "Channel", "Reach", "Catchment", "Cells", "Cell width (m)", "Cell height (m)", "Length L (m)", "Existing structure", "Overall width G (m)", "Concrete (m³)", "Reinforcement (t)", "Blinding (m³)", "Waterproofing (m²)", "Pavement area (m²)", "Aggregate base (m³)", "Bituminous base (m³)", "Wearing course (m³)", "Prime coat (m²)", "Tack coat (m²)", "Marking (m²)",
         "Guardrail (m)", "Posts (no.)", "Compaction (m²)", "Removal of existing (m³)",
         "Concrete cost", "Reinforcement cost", "Blinding cost", "Waterproofing cost", "Aggregate base cost", "Bituminous base cost", "Wearing course cost", "Prime coat cost", "Tack coat cost", "Marking cost", "Guardrail cost", "Posts cost", "Compaction cost", "Removal cost", "Culvert cost (OMR)"]
header(K, kcols, [5, 16, 12, 12, 14, 6, 8, 8, 8, 11, 10] + [11] * 15 + [12] * 15, 52); KL = {name: L(i + 1) for i, name in enumerate(kcols)}
def kr(name, row): return f"{KL[name]}{row}"
for _, r in p.iterrows():
    row = K.max_row + 1; N = kr("Cells", row); B = kr("Cell width (m)", row); Dh = kr("Cell height (m)", row); Lc = kr("Length L (m)", row); G = kr("Overall width G (m)", row); PC = kr("Pavement area (m²)", row); EX = kr("Existing structure", row)
    K.append([int(r.id), r.road, r.channel, r.reach, catchment(r.channel), int(r.cells), float(r.cell_width_m), float(r.cell_height_m), float(r.length_m), (str(r.existing_type) if str(r.existing_type) not in ("", "nan") else "None"),
              f"={N}*{B}+({N}+1)*{const('wall')}", f"=({N}*(({B}+{const('wall')})*({Dh}+0.6)-{B}*{Dh})+0.8)*{Lc}", f"=0.12*{kr('Concrete (m³)', row)}", f"={G}*{Lc}*0.1*1.1", f"=2*{Dh}*{Lc}+{B}*{Lc}", f"={G}*{Lc}*1.1", f"={PC}*0.1", f"={PC}*0.05", f"={PC}*0.05", f"={PC}", f"=2*{PC}", f"={G}*0.1*1.1",
              f"=2*{G}*1.1", f"=ROUNDUP(2*{G}*1.1/4,0)", f"=({Lc}+3)*{G}*1.1", f'=IF({EX}="None",0,{G}*{Lc}*({Dh}+0.5)*1.1)',
              f"={kr('Concrete (m³)', row)}*{rate('conc')}", f"={kr('Reinforcement (t)', row)}*{rate('rebar')}", f"={kr('Blinding (m³)', row)}*{rate('blind')}", f"={kr('Waterproofing (m²)', row)}*{rate('wp')}", f"={kr('Aggregate base (m³)', row)}*{rate('abc')}", f"={kr('Bituminous base (m³)', row)}*{rate('bbc')}", f"={kr('Wearing course (m³)', row)}*{rate('wbc')}",
              f"={kr('Prime coat (m²)', row)}*{rate('pc')}", f"={kr('Tack coat (m²)', row)}*{rate('tc')}", f"={kr('Marking (m²)', row)}*{rate('mark')}", f"={kr('Guardrail (m)', row)}*{rate('rail')}", f"={kr('Posts (no.)', row)}*{rate('post')}", f"={kr('Compaction (m²)', row)}*{rate('cng')}", f"={kr('Removal of existing (m³)', row)}*{rate('removal')}",
              f"=SUM({kr('Concrete cost', row)}:{kr('Removal cost', row)})"])
kl = K.max_row; K.append(["Total", "", "", "", ""] + [f"=SUM({L(j)}2:{L(j)}{kl})" for j in range(6, len(kcols) + 1)])
K.cell(row=K.max_row, column=kcols.index("Existing structure") + 1, value=""); K.cell(row=K.max_row, column=kcols.index("Cell width (m)") + 1, value=""); K.cell(row=K.max_row, column=kcols.index("Cell height (m)") + 1, value="")
for c in K[K.max_row]: c.font = BOLD
for row in K.iter_rows(min_row=2):
    for c in row:
        if c.font != BOLD: c.font = FONT
        if c.column >= 11: c.number_format = "#,##0"
K.freeze_panes = "F2"
# ---- Summary: the totals and the by-catchment block, all by formula
S = wb.create_sheet("Summary", 0)
S["A1"] = "Block 04 - bill of quantities and cost estimate of the channel scheme (200-year design)"; S["A1"].font = Font(name="Arial", size=12, bold=True)
header(S, ["Item", "Quantity", "Unit", "Cost (OMR)", "Cost (million OMR)"], [46, 14, 8, 16, 14]); S.row_dimensions[2].height = 20
tot = C.max_row; dtot = D.max_row; ktot = K.max_row
lines = [("Channels - clearing and grubbing", f"=Channels!{CL['Clearing (m²)']}{tot}", "m²", f"=Channels!{CL['Clearing cost (OMR)']}{tot}"), ("Channels - excavation", f"=Channels!{CL['Excavation (m³)']}{tot}", "m³", f"=Channels!{CL['Excavation cost (OMR)']}{tot}"),
         ("Channels - fill", f"=Channels!{CL['Fill (m³)']}{tot}", "m³", f"=Channels!{CL['Fill cost (OMR)']}{tot}"), ("Channels - riprap", f"=Channels!{CL['Riprap volume (m³)']}{tot}", "m³", f"=Channels!{CL['Riprap cost (OMR)']}{tot}"),
         ("Channels - construction", f"=Channels!{CL['Length (m)']}{tot}", "m", f"=Channels!{CL['Construction cost (OMR)']}{tot}"), ("Channels - land of the corridors (expropriation)", f"=Channels!{CL['Corridor (ha)']}{tot}", "ha", f"=Channels!{CL['Land cost (OMR)']}{tot}"),
         ("Dykes", f"=Dykes!{DL['Length (m)']}{dtot}", "m", f"=Dykes!{DL['Dyke cost (OMR)']}{dtot}"), ("Culverts", f"=Culverts!{KL['Cells']}{ktot}", "cells", f"=Culverts!{KL['Culvert cost (OMR)']}{ktot}")]
for name, qf, unit, cf in lines:
    S.append([name, qf, unit, cf, f"=D{S.max_row + 1}/1000000"])
S.append(["Total", "", "", f"=D{S.max_row - 2}+D{S.max_row - 1}+D{S.max_row}+D{S.max_row - 3}", f"=D{S.max_row + 1}/1000000"])
for c in S[S.max_row]: c.font = BOLD
for row in S.iter_rows(min_row=3):
    for c in row:
        if c.font != BOLD: c.font = FONT
    row[1].number_format = "#,##0"; row[3].number_format = "#,##0"; row[4].number_format = "#,##0.00"
S.append([]); S.append(["By catchment (million OMR)", "Channels, construction", "Channels, land", "Dykes", "Culverts", "Total"])
for c in S[S.max_row]: c.fill = HDR; c.font = HFONT; c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
S.row_dimensions[S.max_row].height = 30; b0 = S.max_row + 1
cc = f"Channels!${CL['Catchment']}$2:${CL['Catchment']}${last}"; dc = f"Dykes!${DL['Catchment']}$2:${DL['Catchment']}${dl}"; kc_ = f"Culverts!${KL['Catchment']}$2:${KL['Catchment']}${kl}"
for g in ("HAW", "SHA", "SAR", "SAK"):
    rw = S.max_row + 1
    S.append([GCA[g], f"=SUMIF({cc},A{rw},Channels!${CL['Construction cost (OMR)']}$2:${CL['Construction cost (OMR)']}${last})/1000000", f"=SUMIF({cc},A{rw},Channels!${CL['Land cost (OMR)']}$2:${CL['Land cost (OMR)']}${last})/1000000",
              f"=SUMIF({dc},A{rw},Dykes!${DL['Dyke cost (OMR)']}$2:${DL['Dyke cost (OMR)']}${dl})/1000000", f"=SUMIF({kc_},A{rw},Culverts!${KL['Culvert cost (OMR)']}$2:${KL['Culvert cost (OMR)']}${kl})/1000000", f"=SUM(B{rw}:E{rw})"])
b1 = S.max_row; S.append(["Total", f"=SUM(B{b0}:B{b1})", f"=SUM(C{b0}:C{b1})", f"=SUM(D{b0}:D{b1})", f"=SUM(E{b0}:E{b1})", f"=SUM(F{b0}:F{b1})"])
for row in S.iter_rows(min_row=b0, max_row=S.max_row):
    for c in row:
        c.font = BOLD if row[0].value == "Total" else FONT
        if c.column >= 2: c.number_format = "#,##0.00"
S.column_dimensions["F"].width = 12
S.append([]); S.append(["The total is the construction of the channels, the land of their corridors, the dykes and the culverts. Quantities: channels measured on the designed sections against the 0.5 m terrain model; dykes on the typical fill section at the adopted height; culverts by the concrete, reinforcement, blinding, waterproofing and road works over them. Rates: sheet Rates (master plan estimate)."])
S[S.max_row][0].font = Font(name="Arial", size=9, italic=True)
path = os.path.join(OUT, "Block 04 - BOQ and cost estimate.xlsx"); wb.save(path); print("written", path, "| channels", len(q), "dykes", len(d), "culverts", len(p))
