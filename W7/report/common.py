"""Paths, figure helpers and the table/figure cross-reference registry shared by the report sections (R1).

Cross-references: a caption made with a key (tcap / fig / chart / map_fig / image_grid ... key="x") registers its number;
T("x") and F("x") in the text give "Table n" / "Figure n". The numbers of the previous pass are read from _refs.json, so
build.py runs the build twice when a number moved (the second pass is always consistent)."""
import os, sys, json
HERE = os.path.dirname(os.path.abspath(__file__)); W7 = os.path.dirname(HERE); CB04 = os.path.dirname(W7); B4 = os.path.dirname(CB04); W3 = os.path.join(CB04, "W3")
sys.path.insert(0, os.path.join(HERE, "lib")); sys.path.insert(0, HERE)
import doc as D  # noqa: E402
import notes as N  # noqa: E402
import data_facts as A  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

IMG = os.path.join(HERE, "img"); MAPS = os.path.join(W7, "maps", "png_r4"); SEDF = os.path.join(W3, "sediment", "figures"); CLIF = os.path.join(W3, "climate", "figures")
HMS_SHOT = next((p for p in (os.path.join(CB04, "Data", "images", "HMS model.png"), os.path.join(B4, "Hydrology", "Claude Block 04", "Data", "images", "HMS model.png")) if os.path.exists(p)), None)

REF_PATH = os.path.join(HERE, "_refs.json")
try: REF = json.load(open(REF_PATH, encoding="utf-8"))
except Exception: REF = {}
NEW = {}
def T(key): return f"Table {REF.get('tab:' + key, '?')}"
def F(key): return f"Figure {REF.get('fig:' + key, '?')}"
def _reg(kind, key, n):
    if key: NEW[f"{kind}:{key}"] = n
    return n
def refs_stable():
    """Write this pass's numbers; True when they equal the numbers the text was built with."""
    json.dump(NEW, open(REF_PATH, "w", encoding="utf-8"), indent=1); stable = (NEW == REF); REF.clear(); REF.update(NEW); NEW.clear(); return stable

def tcap(d, text, key=None): return _reg("tab", key, D.tab_caption(d, text))

def fig(d, path, caption, width=15.5, key=None):
    if not os.path.exists(path): print("   missing figure:", path); return None
    D.picture(d, path, width); return _reg("fig", key, D.fig_caption(d, caption.rstrip(".")))

def chart(d, name, caption, width=15.5, key=None):
    return fig(d, os.path.join(IMG, name + ".png"), caption, width, key=key)

def map_fig(d, key_png, caption, size="A4", key=None):
    path = os.path.join(MAPS, key_png + ".png")
    if not os.path.exists(path): print("   missing map:", path); return None
    return _reg("fig", key, D.wide_figure(d, path, caption.rstrip("."), size=size))

def image_grid(d, items, caption, cols=2, width_cm=7.6, as_table=True, key=None):
    """Several small figures in a grid, each with its own title cell, under ONE caption:
    items = [(path, title), ...]. The caption is a table caption when as_table, else a figure caption."""
    items = [(p, t) for p, t in items if os.path.exists(p)]
    if not items: return None
    if as_table: n = _reg("tab", key, D.tab_caption(d, caption))
    rows = (len(items) + cols - 1) // cols
    t = d.add_table(rows=rows * 2, cols=cols); t.alignment = 1
    tblpr = t._tbl.tblPr
    tblpr.append(parse_xml('<w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:top w:val="nil"/><w:left w:val="nil"/><w:bottom w:val="nil"/><w:right w:val="nil"/><w:insideH w:val="nil"/><w:insideV w:val="nil"/></w:tblBorders>'))
    for k, (path, title) in enumerate(items):
        r, c = (k // cols) * 2, k % cols
        cell = t.rows[r].cells[c]; par = cell.paragraphs[0]; par.alignment = WD_ALIGN_PARAGRAPH.CENTER; run = par.add_run(title); run.bold = True; run.font.size = Pt(8.5)
        par.paragraph_format.space_after = Pt(1); par.paragraph_format.keep_with_next = True
        cell2 = t.rows[r + 1].cells[c]; par2 = cell2.paragraphs[0]; par2.alignment = WD_ALIGN_PARAGRAPH.CENTER; par2.add_run().add_picture(path, width=Cm(width_cm)); par2.paragraph_format.space_after = Pt(6)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(parse_xml('<w:cantSplit xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))
    if not as_table: n = _reg("fig", key, D.fig_caption(d, caption.rstrip(".")))
    return n

def h_app(d, text):
    """A level-1 heading without the template's automatic number (appendices), still in the contents."""
    par = D.h(d, 1, text, page_break=True); ppr = par._p.get_or_add_pPr()
    for old in ppr.findall(qn("w:numPr")): ppr.remove(old)
    ppr.insert(0, parse_xml('<w:numPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:ilvl w:val="0"/><w:numId w:val="0"/></w:numPr>'))
    return par

def merge_first_column_pairs(t, start_row=1):
    """Merge the dam-name cell with the one below it (In / Out rows of the same dam), comment 28."""
    from docx.enum.table import WD_ALIGN_VERTICAL
    rows = t.rows
    for i in range(start_row, len(rows) - 1, 2):
        a, b = rows[i].cells[0], rows[i + 1].cells[0]
        txt = a.text; m = a.merge(b); m.text = ""; par = m.paragraphs[0]; par.text = ""; r = par.add_run(txt); r.bold = True; r.font.size = Pt(6.5)
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT; m.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        par.paragraph_format.keep_with_next = (i + 1 < len(rows) - 1); par.paragraph_format.space_after = Pt(0)   # the merge made a fresh paragraph: keep the table on one page

def n0(v): return f"{v:,.0f}"
def n1(v): return f"{v:,.1f}"
def n2(v): return f"{v:,.2f}"
def n3(v): return f"{v:,.3f}"
lab = A.label
