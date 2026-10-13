"""Build template/Block04_template.docx from the Block 6 channel design report (the project's own report format):
the cover page (texts re-pointed to Block 04, Rev 00, the build month), the header and footer of the body sections, the
styles (Aptos body, blue headings, Caption, toc, table of figures). Everything after the cover is removed.
usage: python make_template.py
"""
import os, copy
from docx import Document
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = r"W:\2224 Flood Protection Scheme in North & South Batinah_MAFWWR\06-WORKING\04-Hydraulic\05 Small Tasks\18 Channel Design\Hydraulic\Block 6\Report\W3\docx\Channel_Design_Block6_Final.docx"
OUT = os.path.join(HERE, "template", "Block04_template.docx"); os.makedirs(os.path.dirname(OUT), exist_ok=True)
REPL = {"Block No. 6": "Block No. 4", "Block 6": "Block 04", "Block No 6": "Block No. 4", "April 2026": "October 2026", "Rev 00": "Rev 00"}

def fix_text(root):
    for t in root.iter(qn("w:t")):
        s = t.text or ""
        for a, b in REPL.items():
            if a in s: s = s.replace(a, b)
        t.text = s

d = Document(SRC); body = d.element.body
kids = [k for k in body if k.tag != qn("w:sectPr")]
# the cover is everything up to and including the first paragraph that carries a section break
end = None
for i, k in enumerate(kids):
    if k.tag == qn("w:p") and k.find(qn("w:pPr")) is not None and k.find(qn("w:pPr")).find(qn("w:sectPr")) is not None:
        end = i; break
assert end is not None, "no section break after the cover"
for k in kids[end + 1:]: body.remove(k)
# drop leading empties: the two blank paragraphs above the emblem pushed the revision lines onto a second page
for k in kids[:2]:
    if k.tag == qn('w:p') and not ''.join(t.text or '' for t in k.iter(qn('w:t'))).strip(): body.remove(k)
fix_text(body)
for s in d.sections:
    for hf in (s.header, s.footer, s.first_page_header, s.first_page_footer):
        if hf._element is not None: fix_text(hf._element)
# the final section properties (body pages: Letter portrait, 2.5 cm margins, the running header) stay as the trailing sectPr
# heading numbers in the heading look (comment 4): the numbering levels linked to Heading 1-3 take the heading colour, bold, size
from docx.oxml import parse_xml
num = d.part.numbering_part.element
styles = {s_.get(qn('w:styleId')): s_ for s_ in d.styles.element.findall(qn('w:style'))}
for sid, size in (('Heading1', 28), ('Heading2', 24), ('Heading3', 24)):
    st = styles.get(sid)
    if st is None: continue
    ppr = st.find(qn('w:pPr')); npr = ppr.find(qn('w:numPr')) if ppr is not None else None
    if npr is None: continue
    num_id = npr.find(qn('w:numId')).get(qn('w:val')); ilvl = (npr.find(qn('w:ilvl')).get(qn('w:val')) if npr.find(qn('w:ilvl')) is not None else '0')
    n_el = [x for x in num.findall(qn('w:num')) if x.get(qn('w:numId')) == num_id]
    if not n_el: continue
    abs_id = n_el[0].find(qn('w:abstractNumId')).get(qn('w:val'))
    for an in num.findall(qn('w:abstractNum')):
        if an.get(qn('w:abstractNumId')) != abs_id: continue
        for lvl in an.findall(qn('w:lvl')):
            if lvl.get(qn('w:ilvl')) == ilvl:
                for old in lvl.findall(qn('w:rPr')): lvl.remove(old)
                lvl.append(parse_xml(f'<w:rPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:rFonts w:asciiTheme="minorHAnsi" w:hAnsiTheme="minorHAnsi"/><w:b/><w:color w:val="0000CC"/><w:sz w:val="{size}"/><w:szCs w:val="{size}"/></w:rPr>'))
                print('numbering level styled:', sid, 'num', num_id, 'lvl', ilvl)
d.save(OUT)
d2 = Document(OUT); print("template written:", OUT, "| sections:", len(d2.sections), "| cover paragraphs:", [q.text[:40] for q in d2.paragraphs if q.text.strip()][:12])
print("body section:", f"{d2.sections[-1].page_width.cm:.1f}x{d2.sections[-1].page_height.cm:.1f} cm", "| header:", [q.text[:70] for q in d2.sections[-1].header.paragraphs if q.text.strip()])
