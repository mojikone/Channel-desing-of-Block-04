"""Step 05 - peak flow and volume at every element, Block 6 layout ('peak_flow_volume_all_elements.xlsx'), for Basin 3 and Basin 2.

Sheets 'Peak Flow m3s' and 'Volume Mm3': element | type | 2yr ... PMP.  Values are the depth-area results at each element's
own analysis point (subbasin outflow, reach outflow, junction combined flow, reservoir outflow); volume = sum x 300 s / 1e6.
Reads  ..\..\results\basin{2,3}_peak_flow_volume_long.csv (extract_b2.py, 2026-10-05)
Writes W3/elements/xlsx/peak_flow_volume_all_elements_B3.xlsx and _B2.xlsx (+ the long CSV copies)
"""
import sys, re, shutil
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as L
sys.path.insert(0, str(Path(__file__).resolve().parent))
from config import RPS, RESULTS_CB04, W3, ensure

OUT = ensure(W3 / "elements" / "xlsx"); FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F4E79"); thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
def skey(n):
    m = re.match(r"^([A-Z]{3})-([SJRO])(\d+)$", n); return (0, m.group(1), "SJRO".index(m.group(2)), int(m.group(3))) if m else (1, n, 0, 0)
for b in ("3", "2"):
    src = RESULTS_CB04 / f"basin{b}_peak_flow_volume_long.csv"; df = pd.read_csv(src); shutil.copy2(src, OUT / f"peak_flow_volume_B{b}_long.csv")
    kind = df.drop_duplicates("element").set_index("element")["type"]
    wb = Workbook(); wb.remove(wb.active)
    for sheet, val, fmt in (("Peak Flow m3s", "peak_m3s", "0.0"), ("Volume Mm3", "vol_Mm3", "0.000")):
        pv = df.pivot(index="element", columns="rp", values=val)[RPS]; pv = pv.loc[sorted(pv.index, key=skey)]
        ws = wb.create_sheet(sheet)
        for j, h in enumerate(["element", "type"] + RPS, 1):
            c = ws.cell(row=1, column=j, value=h); c.font = Font(name=FONT, bold=True, color="FFFFFF"); c.fill = HDR_FILL; c.alignment = Alignment(horizontal="center"); c.border = BOX
        for i, (el, row) in enumerate(pv.iterrows(), 2):
            ws.cell(row=i, column=1, value=el).font = Font(name=FONT, size=10, bold=True); ws.cell(row=i, column=2, value=kind[el]).font = Font(name=FONT, size=10)
            for j, rp in enumerate(RPS, 3):
                c = ws.cell(row=i, column=j, value=round(float(row[rp]), 4)); c.number_format = fmt; c.font = Font(name=FONT, size=10); c.border = BOX
        ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 11
        for j in range(3, 3 + len(RPS)): ws.column_dimensions[L(j)].width = 10
        ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:{L(2 + len(RPS))}{1 + len(pv)}"
    rd = wb.create_sheet("Read me")
    lines = [f"Peak flow (m³/s) and volume (Mm³) at every element of Basin {b} - HEC-HMS 4.12 project HEC_HMS_Block_04, depth-area analyses 'DA <rp> B{b}' (2026-10-05).",
             "Basin 2 = dam catchments (Basin 1 elements) + channel network, no reservoirs. Basin 3 = Basin 2 + the eight dam reservoirs; channel reach sections sized on the reservoir-routed 200-year flow." if b == "3" else
             "Basin 2 = dam catchments (Basin 1 elements) + channel network, NO reservoirs (the undammed case); channel reach sections sized on its own 200-year flow.",
             "Each value is read at the element's own analysis point (the depth-area analysis reduces the storm for that element's own catchment area): subbasin outflow, reach outflow, junction combined flow, reservoir outflow.",
             "Volume = sum of the 5-minute ordinates x 300 s / 1e6. Element names: <global catchment>-S/R/J/O<n> (HAW Al Hawasinah, SHA Shafan, SAR Sarami, SAK Sakhin); O = outfall to the sea; dam junctions and reservoirs keep the dam name.",
             "Muskingum-Cunge channel widths are the first sizing pass (n 0.1, 2 m depth at Q200); a second pass would change a few reaches by 10-35 % (see Claude Block 04/README.md)."]
    for i, t in enumerate(lines, 1): rd.cell(row=i, column=1, value=t).font = Font(name=FONT, size=10)
    rd.column_dimensions["A"].width = 170
    out = OUT / f"peak_flow_volume_all_elements_B{b}.xlsx"; wb.save(out); print("written", out, "| elements:", pv.shape[0])
