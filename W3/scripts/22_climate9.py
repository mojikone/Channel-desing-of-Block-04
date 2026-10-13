"""Step 22 - climate at each of the nine dams (replaces step 12 for the report): the 9 km reanalysis record of step 17
(daily 2001-2023, hourly wind 2015-2024) and the MRMWR gauges of step 18.
Figures : climate/figures/windrose_<dam>.png (one per dam), windroses_all.png (3 x 3), monthly_all.png (3 x 3 temperature
          and rainfall), annual_rain_all.png (3 x 3 annual rainfall)
Summary : climate/climate9_summary.json (per dam: temperature, rainfall, wind, cell; gauges)
Workbook: climate/Block04_Climate.xlsx (one sheet per dam: daily data + monthly and annual formulas + wind sector table;
          Summary and Gauges sheets)
"""
import sys, json, math
from pathlib import Path
import numpy as np, pandas as pd
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as L
sys.path.insert(0, str(Path(__file__).resolve().parent))
from config import DAMS9, label, W3, ensure
import figstyle as FS; FS.apply()

CL = W3 / "climate"; FIG = ensure(CL / "figures"); DATA = CL / "data_om"; MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
SECT = ["N", "NNE", "NE", "ENE", "E", "ESE", "SE", "SSE", "S", "SSW", "SW", "WSW", "W", "WNW", "NW", "NNW"]; CLASSES = ["0-2", "2-4", "4-6", "6-8", ">8"]; EDGES = [0, 2, 4, 6, 8, 99]
SPEED_COL = ["#C6D6EA", "#8FADD1", "#4F81BD", "#1F497D", "#C0504D"]; names = {"0-2": "under 2 m/s", "2-4": "2 to 4", "4-6": "4 to 6", "6-8": "6 to 8", ">8": "over 8 m/s"}
cells = pd.read_csv(DATA / "cells.csv").set_index("dam")
near = pd.read_csv(CL / "gauges" / "gauges_near_dams.csv"); gst = pd.read_csv(CL / "gauges" / "gauge_annual.csv", index_col=0); gann = pd.read_csv(CL / "gauges" / "gauge_annual_series.csv")

def load(dam):
    d = DATA / dam.replace(" ", "_"); D = pd.read_csv(d / "daily_2001_2023.csv", parse_dates=["date"]).dropna(); D["year"] = D.date.dt.year; D["month"] = D.date.dt.month
    W = pd.read_csv(d / "hourly_wind_2015_2024.csv", parse_dates=["datetime_local"]).dropna(); W["sector"] = ((W.wd10m_deg % 360 + 11.25) // 22.5).astype(int) % 16; W["cls"] = pd.cut(W.ws10m_ms, EDGES, labels=CLASSES, right=False, include_lowest=True)
    return D, W

def rose_table(W):
    t = pd.crosstab(W.sector, W.cls).reindex(index=range(16), columns=CLASSES, fill_value=0) / len(W) * 100; return t

def draw_rose(ax, t, mean, calm, title, rmax):
    th = np.deg2rad(np.arange(16) * 22.5); wdt = 2 * np.pi / 16 * 0.86; bottom = np.zeros(16)
    for c, col in zip(CLASSES, SPEED_COL):
        v = t[c].values; ax.bar(th, v, width=wdt, bottom=bottom, color=col, edgecolor="white", linewidth=0.3, zorder=3); bottom += v
    ax.set_theta_zero_location("N"); ax.set_theta_direction(-1); ax.set_xticks(np.deg2rad(np.arange(0, 360, 45))); ax.set_xticklabels(["N", "NE", "E", "SE", "S", "SW", "W", "NW"], fontsize=7.2, color="#404040")
    ticks = list(range(4, int(rmax) + 1, 4)); ax.set_ylim(0, rmax); ax.set_yticks(ticks); ax.set_yticklabels([f"{t_} %" for t_ in ticks], fontsize=5.8, color="#7A7A7A"); ax.set_rlabel_position(78)
    ax.grid(color="#E6E6E6", lw=0.5); ax.spines["polar"].set_color("#BFBFBF"); ax.set_title(f"{title}\nmean {mean:.1f} m/s, calm {calm:.0f} % of hours", fontsize=7.6, color="#404040", pad=8)

summary = {}; Ds = {}; Ws = {}; tables = {}
for dam in DAMS9:
    D, W = load(dam); Ds[dam] = D; Ws[dam] = W; t = rose_table(W); tables[dam] = t
    ny = D.year.nunique(); mon = D.groupby("month").agg(t_mean=("t_mean_c", "mean"), t_max=("t_max_c", "mean"), t_min=("t_min_c", "mean"), rain=("rain_mm", "sum")); mon["rain"] /= ny
    ann = D.groupby("year").agg(rain=("rain_mm", "sum"), t_mean=("t_mean_c", "mean")); tot = t.sum(axis=1); top = tot.sort_values(ascending=False)
    ng = near[near.dam == dam].merge(gst[["years", "mean_annual", "first", "last", "Local_Name"]], left_on="site_id", right_index=True).sort_values("dist_km")
    ng = ng[ng.years >= 4]; gg = ng.iloc[0] if len(ng) else None
    summary[dam] = dict(label=label(dam), cell=[float(cells.loc[dam, "cell_lat"]), float(cells.loc[dam, "cell_lon"])], t_mean=round(float(D.t_mean_c.mean()), 1), t_max_month=MONTHS[int(mon.t_max.idxmax()) - 1], t_max_value=round(float(mon.t_max.max()), 1),
                        t_min_month=MONTHS[int(mon.t_min.idxmin()) - 1], t_min_value=round(float(mon.t_min.min()), 1), t_abs_max=round(float(D.t_max_c.max()), 1), t_abs_min=round(float(D.t_min_c.min()), 1),
                        rain_mean_annual=round(float(ann.rain.mean()), 0), rain_min=round(float(ann.rain.min()), 0), rain_max=round(float(ann.rain.max()), 0), wet_days=round(float((D.rain_mm >= 1).sum() / ny), 1),
                        rain_share_nov_apr=round(float(mon.rain[[11, 12, 1, 2, 3, 4]].sum() / mon.rain.sum() * 100), 0), t_trend_c_per_decade=round(float(np.polyfit(ann.index.values, ann.t_mean.values, 1)[0] * 10), 2),
                        wind_mean=round(float(W.ws10m_ms.mean()), 1), wind_calm_pct=round(float((W.ws10m_ms < 1).mean() * 100), 0), wind_top=[(SECT[i], round(float(top.loc[i]), 1)) for i in top.index[:3]],
                        gauge=(dict(name=str(gg["name"]), site=str(gg.site_id), dist_km=float(gg.dist_km), years=int(gg.years), first=int(gg["first"]), last=int(gg["last"]), mean_annual=float(gg.mean_annual)) if gg is not None else None),
                        monthly=mon.round(1).to_dict(orient="index"), annual=ann.round(1).to_dict(orient="index"))
json.dump(summary, open(CL / "climate9_summary.json", "w"), indent=1, default=float)
groups = {}
for dam in DAMS9: groups.setdefault(tuple(summary[dam]["cell"]), []).append(dam)
print("cells:", {k: v for k, v in groups.items()})
# ---- figures
rmax = 4 * math.ceil(max(tables[d].sum(axis=1).max() for d in DAMS9) / 4.0)
for dam in DAMS9:
    fig = plt.figure(figsize=(4.2, 4.6)); ax = fig.add_subplot(111, projection="polar"); draw_rose(ax, tables[dam], summary[dam]["wind_mean"], summary[dam]["wind_calm_pct"], f"{label(dam)} - wind at 10 m, 2015-2024", rmax)
    fig.legend(handles=[Patch(color=c, label=names[k]) for k, c in zip(CLASSES, SPEED_COL)], frameon=False, fontsize=6.8, loc="lower center", ncol=3, title="a bar points to where the wind comes from", title_fontsize=6.8)
    fig.tight_layout(rect=(0, 0.1, 1, 1)); fig.savefig(FIG / f"windrose_{dam.replace(' ', '_')}.png", dpi=200); plt.close(fig)
fig, axes = plt.subplots(3, 3, figsize=(7.6, 8.6), subplot_kw={"projection": "polar"})
for ax, dam in zip(axes.ravel(), DAMS9): draw_rose(ax, tables[dam], summary[dam]["wind_mean"], summary[dam]["wind_calm_pct"], label(dam, True), rmax)
fig.legend(handles=[Patch(color=c, label=names[k]) for k, c in zip(CLASSES, SPEED_COL)], frameon=False, fontsize=7.4, loc="lower center", ncol=5, title="wind speed at 10 m, all hours 2015-2024; a bar points to where the wind comes from", title_fontsize=7.4)
fig.tight_layout(rect=(0, 0.05, 1, 1), h_pad=1.6); fig.savefig(FIG / "windroses_all.png", dpi=200); plt.close(fig)
fig, axes = plt.subplots(3, 3, figsize=(7.6, 7.4), sharex=True)
x = np.arange(12)
for ax, dam in zip(axes.ravel(), DAMS9):
    m = pd.DataFrame(summary[dam]["monthly"]).T; m.index = m.index.astype(int); m = m.sort_index()
    ax.bar(x, m.rain, width=0.6, color=FS.PALE, label="rainfall, mm/month"); ax.set_ylim(0, 40); FS.finish(ax)
    ax2 = ax.twinx(); ax2.plot(x, m.t_max, color=FS.RED, lw=1.1, label="mean daily max"); ax2.plot(x, m.t_mean, color=FS.GREY, lw=1.1, label="daily mean"); ax2.plot(x, m.t_min, color=FS.BLUE, lw=1.1, label="mean daily min"); ax2.set_ylim(0, 50)
    for sp in ("top",): ax2.spines[sp].set_visible(False)
    ax2.tick_params(labelsize=6.5, colors=FS.MUTED); ax.tick_params(labelsize=6.5); ax.set_title(f"{label(dam, True)}: {summary[dam]['rain_mean_annual']:.0f} mm/yr, {summary[dam]['t_mean']:.1f} °C", fontsize=7.4, loc="left")
    ax.set_xticks(x); ax.set_xticklabels([mm[0] for mm in MONTHS], fontsize=6.2)
    if dam == DAMS9[0]: h1, l1 = ax.get_legend_handles_labels(); h2, l2 = ax2.get_legend_handles_labels()
for ax in axes[:, 0]: ax.set_ylabel("rain, mm", fontsize=7)
fig.legend(h1 + h2, l1 + l2, frameon=False, fontsize=7.2, loc="lower center", ncol=4); fig.suptitle("Monthly rainfall (bars, left axis) and air temperature (lines, right axis, °C) at the dams, 2001-2023", fontsize=8.5, color=FS.INK)
fig.tight_layout(rect=(0, 0.04, 1, 0.97), h_pad=1.4); fig.savefig(FIG / "monthly_all.png", dpi=200); plt.close(fig)
fig, axes = plt.subplots(3, 3, figsize=(7.6, 6.8), sharex=True, sharey=True)
for ax, dam in zip(axes.ravel(), DAMS9):
    a = pd.DataFrame(summary[dam]["annual"]).T; a.index = a.index.astype(int); a = a.sort_index(); ax.bar(a.index, a.rain, color=FS.PALE, width=0.75); ax.axhline(a.rain.mean(), color=FS.GREY, lw=0.8, ls="--")
    ax.set_title(f"{label(dam, True)}: mean {a.rain.mean():.0f} mm", fontsize=7.4, loc="left"); FS.finish(ax); ax.tick_params(labelsize=6.2); ax.set_xticks(range(2001, 2024, 4))
for ax in axes[:, 0]: ax.set_ylabel("mm/yr", fontsize=7)
fig.suptitle("Annual rainfall at the dams, 2001-2023 (9 km reanalysis)", fontsize=8.5, color=FS.INK); fig.tight_layout(rect=(0, 0, 1, 0.97)); fig.savefig(FIG / "annual_rain_all.png", dpi=200); plt.close(fig)
# ---- workbook
FONT = "Arial"; HDR_FILL = PatternFill("solid", fgColor="1F4E79"); thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F = lambda **k: Font(name=FONT, size=k.pop("size", 10), **k); BLUE_IN = Font(name=FONT, size=10, color="0000FF"); GREYF = Font(name=FONT, size=9, color="7F7F7F")
def hdr(ws, row, labels, widths=None, height=30):
    for j, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=j, value=lab); c.font = F(bold=True, color="FFFFFF"); c.fill = HDR_FILL; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOX
    ws.row_dimensions[row].height = height
    if widths:
        for j, w in enumerate(widths, 1): ws.column_dimensions[L(j)].width = w
wb = Workbook(); rd = wb.active; rd.title = "Read me"
lines = ["Block 04 dams - climate record at each dam", "",
         "Reanalysis: Open-Meteo historical archive, model best_match (ECMWF IFS 9 km from 2017, ERA5 / ERA5-Land before), read at the dam axis: daily 2001-2023 temperature (mean, max, min) and precipitation; hourly 2015-2024 wind speed and direction at 10 m, local time.",
         f"Grid cells (lat, lon): " + "; ".join(f"{', '.join(label(d, True) for d in v)}: {k[0]:.3f}, {k[1]:.3f}" for k, v in groups.items()) + ". Dams in one cell share the record.",
         "Gauges: MRMWR rain gauges within 15 km of each dam, gap-filled daily record (project 2331, July 2024): sheet 'Gauges'. Records are short (2015 on for the telemetry gauges); they are shown next to the reanalysis, not instead of it.",
         "Dam sheets: daily data (blue) with monthly and annual statistics by formula; wind sector x speed-class table in % of hours (values). Hawasinah EX is reported as Hawasinah UP; Sakhin 08 as Sakhin 08 (Mehmum).",
         f"Built by W3/scripts/22_climate9.py on {pd.Timestamp.today():%Y-%m-%d}."]
for i, t in enumerate(lines, 1): rd.cell(row=i, column=1, value=t).font = F(bold=(i == 1), size=12 if i == 1 else 10)
rd.column_dimensions["A"].width = 170
sm = wb.create_sheet("Summary"); hdr(sm, 1, ["Dam", "Mean T (°C)", "Hottest month", "Mean daily max (°C)", "Coolest month", "Mean daily min (°C)", "Absolute max (°C)", "Absolute min (°C)", "Mean annual rain (mm)", "Driest year (mm)", "Wettest year (mm)", "Wet days/yr", "Nov-Apr share (%)", "T trend (°C/decade)", "Mean wind (m/s)", "Calm (%)", "Prevailing sectors", "Nearest gauge", "Gauge km", "Gauge years", "Gauge mean annual (mm)"], [18, 9, 9, 9, 9, 9, 9, 9, 10, 9, 9, 9, 9, 9, 9, 8, 22, 22, 8, 8, 10], 46)
for i, dam in enumerate(DAMS9, 2):
    s = summary[dam]; g = s["gauge"]
    vals = [label(dam), f"='{label(dam, True)}'!N16", s["t_max_month"], f"='{label(dam, True)}'!O16", s["t_min_month"], f"='{label(dam, True)}'!P16", s["t_abs_max"], s["t_abs_min"], f"='{label(dam, True)}'!Q16", s["rain_min"], s["rain_max"], s["wet_days"], s["rain_share_nov_apr"], s["t_trend_c_per_decade"], s["wind_mean"], s["wind_calm_pct"],
            ", ".join(f"{a} {b:.0f} %" for a, b in s["wind_top"]), (g["name"] if g else "-"), (g["dist_km"] if g else "-"), (g["years"] if g else "-"), (g["mean_annual"] if g else "-")]
    for j, v in enumerate(vals, 1): c = sm.cell(row=i, column=j, value=v); c.border = BOX; c.number_format = "0.0"; c.font = F(bold=(j == 1)) if (j == 1 or (isinstance(v, str) and v.startswith("="))) else BLUE_IN
sm.freeze_panes = "B2"
for dam in DAMS9:
    D, W = Ds[dam], Ws[dam]; ws = wb.create_sheet(label(dam, True)); ny = D.year.nunique(); n = len(D); last = 1 + n; q = f"'{label(dam, True)}'"
    hdr(ws, 1, ["Date", "T mean (°C)", "T max (°C)", "T min (°C)", "Rain (mm)", "Year", "Month"], [12, 10, 10, 10, 10, 7, 7])
    for i, r in enumerate(D.itertuples(), 2):
        ws.cell(row=i, column=1, value=r.date.to_pydatetime()).number_format = "yyyy-mm-dd"
        for j, v in enumerate((r.t_mean_c, r.t_max_c, r.t_min_c, r.rain_mm), 2): c = ws.cell(row=i, column=j, value=round(float(v), 2)); c.font = BLUE_IN; c.number_format = "0.0"
        ws.cell(row=i, column=6, value=f"=YEAR(A{i})"); ws.cell(row=i, column=7, value=f"=MONTH(A{i})")
    rng = lambda col: f"${col}$2:${col}${last}"
    for j, lab in enumerate(["Month", "Mean daily T (°C)", "Mean daily max (°C)", "Mean daily min (°C)", "Rainfall (mm/month)"], 13): hdr(ws, 2, [], None); c = ws.cell(row=2, column=j, value=lab); c.font = F(bold=True, color="FFFFFF"); c.fill = HDR_FILL; c.alignment = Alignment(horizontal="center", wrap_text=True); c.border = BOX
    for j in range(13, 18): ws.column_dimensions[L(j)].width = 13
    ws.row_dimensions[2].height = 30
    for m in range(1, 13):
        r = 2 + m; ws.cell(row=r, column=13, value=MONTHS[m - 1]).font = F(bold=True)
        for j, f in enumerate([f"=AVERAGEIFS({rng('B')},{rng('G')},{m})", f"=AVERAGEIFS({rng('C')},{rng('G')},{m})", f"=AVERAGEIFS({rng('D')},{rng('G')},{m})", f"=SUMIFS({rng('E')},{rng('G')},{m})/{ny}"], 14):
            c = ws.cell(row=r, column=j, value=f); c.number_format = "0.0"; c.border = BOX
    ws.cell(row=16, column=13, value="Year").font = F(bold=True)
    for j, f in enumerate([f"=AVERAGE({rng('B')})", f"=MAX(O3:O14)", f"=MIN(P3:P14)", "=SUM(Q3:Q14)"], 14): c = ws.cell(row=16, column=j, value=f); c.number_format = "0.0"; c.border = BOX; c.font = F(bold=True)
    ws.cell(row=17, column=13, value="N16 mean T; O16 hottest month mean max; P16 coolest month mean min; Q16 mean annual rain").font = GREYF
    ws.cell(row=19, column=13, value=f"Wind at 10 m, share of hours (%) by sector (blows FROM) and speed class, {W.datetime_local.dt.year.min()}-{W.datetime_local.dt.year.max()}, {len(W):,} hours").font = F(bold=True)
    for j, lab in enumerate(["Sector"] + [f"{c} m/s" for c in CLASSES] + ["All"], 13): c = ws.cell(row=20, column=j, value=lab); c.font = F(bold=True, color="FFFFFF"); c.fill = HDR_FILL; c.alignment = Alignment(horizontal="center"); c.border = BOX
    t = tables[dam]
    for i, sname in enumerate(SECT):
        r = 21 + i; ws.cell(row=r, column=13, value=sname).font = F(bold=True)
        for j, c_ in enumerate(CLASSES, 14): cell = ws.cell(row=r, column=j, value=round(float(t.loc[i, c_]), 3)); cell.font = BLUE_IN; cell.number_format = "0.0"; cell.border = BOX
        cell = ws.cell(row=r, column=19, value=f"=SUM(N{r}:R{r})"); cell.number_format = "0.0"; cell.border = BOX
    ws.cell(row=37, column=13, value="Mean speed (m/s)").font = F(bold=True); ws.cell(row=37, column=14, value=summary[dam]["wind_mean"]).font = BLUE_IN
    ws.cell(row=38, column=13, value="Calm, below 1 m/s (%)").font = F(bold=True); ws.cell(row=38, column=14, value=summary[dam]["wind_calm_pct"]).font = BLUE_IN
    ws.freeze_panes = "A2"
gs = wb.create_sheet("Gauges"); hdr(gs, 1, ["Dam", "Gauge", "Site ID", "Type", "Elevation (m)", "Distance (km)", "Complete years", "First", "Last", "Mean annual (mm)", "Median (mm)", "Max (mm)", "Min (mm)"], [18, 30, 13, 9, 9, 9, 9, 7, 7, 10, 9, 9, 9], 40)
r = 2
for dam in DAMS9:
    for g in near[near.dam == dam].sort_values("dist_km").itertuples():
        st = gst.loc[g.site_id] if g.site_id in gst.index else None
        vals = [label(dam), g.name, g.site_id, g.type, g.elev_m, g.dist_km] + ([int(st.years), int(st["first"]), int(st["last"]), st.mean_annual, st.median_annual, st.max_annual, st.min_annual] if st is not None else ["-"] * 7)
        for j, v in enumerate(vals, 1): c = gs.cell(row=r, column=j, value=v); c.border = BOX; c.font = F(bold=(j == 1)) if j <= 4 else BLUE_IN
        r += 1
gs.cell(row=r + 1, column=1, value="MRMWR gauges, gap-filled daily record (project 2331, 20240714_MRMWR_Rainfall Data_Complete, read-only); complete year = at least 330 days.").font = GREYF
for s_ in wb.worksheets:
    for row in s_.iter_rows():
        for c in row:
            if c.font is None or c.font.name != FONT: c.font = Font(name=FONT, size=c.font.size or 10, bold=c.font.bold, color=c.font.color)
out = CL / "Block04_Climate.xlsx"; wb.save(out); print("written", out, "| figures:", sorted(p.name for p in FIG.glob("*.png"))[:6], "...")
for dam in DAMS9: s = summary[dam]; print(f"{label(dam):20s} T {s['t_mean']} | rain {s['rain_mean_annual']:.0f} mm ({s['rain_min']:.0f}-{s['rain_max']:.0f}) | wind {s['wind_mean']} m/s {s['wind_top'][0]} | gauge {s['gauge']['name'] if s['gauge'] else '-'} {s['gauge']['mean_annual'] if s['gauge'] else ''}")
