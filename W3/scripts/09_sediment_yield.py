"""Step 09 - sediment yield of the eight dam catchments: RUSLE factors (pipeline output) + Williams (1977) sediment
delivery ratio, in the layout of 'B6 DAM RUSLE.xlsx' (Block 6), all derived numbers as Excel formulas.

Inputs : sediment/rusle/Output/CSV/RUSLE_Results.csv   (R, K, LS, C, P, mean soil loss t/ha/yr per catchment)
         sediment/catchments/Dam_catchments.gpkg        (area, weighted CN, subbasin list)
         model/Basin_2.basin + Basin_2.sqlite            (longest flow path to the dam = subbasin longest flowpath + reaches)
         HMS terrain elevation.tif (10 m)                (basin relief = highest cell in the catchment - outlet elevation)
Williams (1977): SDR = 1.366e-11 * A^-0.1 * RR^0.363 * CN^5.444   (A km2, RR = relief / basin length, CN weighted)
Sediment delivered (m3/km2/yr) = soil loss (t/ha/yr) * 100 / density * SDR, density 1.5 t/m3 as in Block 6.
Writes sediment/B4_DAM_RUSLE.xlsx and sediment/sediment_yield.csv (input of step 11, reservoir sedimentation).
"""
import sys, sqlite3
from pathlib import Path
import numpy as np, pandas as pd, geopandas as gpd, rasterio
from rasterio import mask as rmask
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as L
sys.path.insert(0, str(Path(__file__).resolve().parent))
from config import DAMS, MODEL, HMS, W3, ensure
import basinfile as bf

SED = ensure(W3 / "sediment"); RES = SED / "rusle" / "Output" / "CSV" / "RUSLE_Results.csv"
DEM = HMS / "terrain" / "NSA_5m_B04_Res10m" / "00" / "elevation.tif"
DENSITY = 1.5   # t/m3, Block 6 convention
cat = gpd.read_file(SED / "catchments" / "Dam_catchments.gpkg").set_index("name")
els = bf.parse(MODEL / "Basin_2.basin"); ds = bf.downstream(els)
con = sqlite3.connect(MODEL / "Basin_2.sqlite"); lfp = {n: float(v) for n, v in con.execute("select subbasin_name, longest_length from subbasin_statistics")}; con.close()
rus = pd.read_csv(RES); rus["Name"] = rus["Name"].astype(str); rus = rus.set_index("Name")
print("RUSLE columns:", list(rus.columns))

def flow_path_km(dam):
    """Longest flow path to the dam: for each upstream subbasin, its longest flowpath + the reaches between its outlet and the dam."""
    best = 0.0
    for s in cat.loc[dam, "subbasins"].split(", "):
        length = lfp.get(s, 0.0); cur = ds.get(s); seen = set()
        while cur and cur != dam and cur not in seen:
            seen.add(cur)
            if els[cur]["type"] == "Reach": length += bf.num(els[cur], "Length", 0.0)
            cur = ds.get(cur)
        best = max(best, length)
    return best / 1000.0

rows = []
with rasterio.open(DEM) as src:
    nod = src.nodata
    for dam in DAMS:
        geom = cat.loc[dam, "geometry"]; arr, _ = rmask.mask(src, [geom], crop=True, nodata=np.nan); z = arr[0]; z = z[np.isfinite(z) & (z < 1e30) & (z > -1e4)]
        jx, jy = cat.loc[dam, "jx"], cat.loc[dam, "jy"]; w = 100
        out_arr, _ = rmask.mask(src, [geom.intersection(gpd.points_from_xy([jx], [jy])[0].buffer(w))], crop=True, nodata=np.nan); zo = out_arr[0]; zo = zo[np.isfinite(zo) & (zo < 1e30) & (zo > -1e4)]
        z_out = float(zo.min()) if zo.size else float(np.nan); z_max = float(z.max()); relief = z_max - z_out; lkm = flow_path_km(dam); rr = relief / (lkm * 1000)
        r = rus.loc[dam]
        rows.append(dict(name=dam, soil_loss_t_ha_yr=float(r["Mean Soil Loss (t/ha/yr)"]), R=float(r["R"]), K=float(r["K"]), LS=float(r["LS"]), C=float(r["C"]), P=float(r["P"]),
                         area_km2=float(cat.loc[dam, "area_km2"]), z_max_m=round(z_max, 1), z_outlet_m=round(z_out, 1), relief_m=round(relief, 1), flow_path_km=round(lkm, 2),
                         relief_ratio=round(rr, 4), cn=float(cat.loc[dam, "cn_w"]), fsl_capacity_Mm3=np.nan))
df = pd.DataFrame(rows)
df["sdr_williams"] = 1.366e-11 * df.area_km2 ** -0.1 * df.relief_ratio ** 0.363 * df.cn ** 5.444
df["yield_m3_km2_yr"] = (df.soil_loss_t_ha_yr * 100 / DENSITY * df.sdr_williams).round(1)
df["yield_m3_yr"] = (df.yield_m3_km2_yr * df.area_km2).round(0)
df.to_csv(SED / "sediment_yield.csv", index=False); pd.set_option("display.width", 220); print(df.drop(columns="fsl_capacity_Mm3").to_string(index=False))

# ---- workbook in the Block 6 layout + the extra columns
FONT = "Arial"; HDR_FILL = PatternFill("solid", fgColor="1F4E79"); thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BLUE = Font(name=FONT, size=10, color="0000FF"); BLK = Font(name=FONT, size=10); GREY = Font(name=FONT, size=9, color="7F7F7F")
wb = Workbook(); ws = wb.active; ws.title = "B4 DAM RUSLE"
cols = ["Name", "Mean Soil Loss (t/ha/yr)", "R", "K", "LS", "C", "P", "Area km2", "Basin Relief Ratio", "CN", "SDR method", "SDR value", "Sediment Density 1.3-1.8 (t/m3)",
        "Mean Soil Loss (m3/km2/year) considering SDR", "Sediment yield (m3/yr)", "Sediment in 50 yr (Mm3)", "Sediment in 100 yr (Mm3)", "Relief (m)", "Longest flow path (km)", "Highest point (m)", "Outlet (m)"]
for j, h in enumerate(cols, 1):
    c = ws.cell(row=1, column=j, value=h); c.font = Font(name=FONT, bold=True, color="FFFFFF"); c.fill = HDR_FILL; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOX
    ws.column_dimensions[L(j)].width = 13 if j > 1 else 15
ws.row_dimensions[1].height = 48; ws.column_dimensions["K"].width = 16
for i, r in enumerate(df.itertuples(), 2):
    vals = [r.name, r.soil_loss_t_ha_yr, r.R, r.K, r.LS, r.C, r.P, r.area_km2, r.relief_ratio, r.cn, "Williams (1977)", f"=1.366*10^-11*H{i}^-0.1*I{i}^0.363*J{i}^5.444",
            DENSITY if i == 2 else "=$M$2", f"=(B{i}*100/M{i})*L{i}", f"=N{i}*H{i}", f"=O{i}*50/1000000", f"=O{i}*100/1000000", r.relief_m, r.flow_path_km, r.z_max_m, r.z_outlet_m]
    fm = ["@", "0.000", "0.000", "0.0000", "0.000", "0.0000", "0.00", "0.00", "0.0000", "0.0", "@", "0.0000", "0.0", "0.0", "#,##0", "0.000", "0.000", "0", "0.00", "0", "0"]
    for j, (v, f) in enumerate(zip(vals, fm), 1):
        c = ws.cell(row=i, column=j, value=v); c.number_format = f; c.border = BOX
        c.font = BLUE if (j in (2, 3, 4, 5, 6, 7, 8, 9, 10, 18, 19, 20, 21) or (j == 13 and i == 2)) else BLK
n = 2 + len(df)
notes = ["Blue = data, black = formulas. RUSLE factors: Scripts of github.com/mojikone/RUSLE run on the eight dam catchments (R from SM2RAIN monthly climatology, K from SoilGrids 0-5 cm (Williams 1995), LS from Copernicus GLO-30 (Moore & Burch), C from Landsat 8 NDVI 2017 (Durigon), P from ESA WorldCover 2020); mean soil loss = zonal mean of R K LS C P at 250 m.",
         "Basin relief ratio = (highest cell of the catchment on the 10 m HMS terrain - lowest cell within 100 m of the dam junction) / longest flow path to the dam (HMS longest flowpath of the farthest subbasin + the reaches down to the dam). CN = area-weighted curve number of the Basin 2 subbasins.",
         "SDR by Williams (1977) as in the Block 6 sheet: SDR = 1.366e-11 A^-0.1 RR^0.363 CN^5.444. Sediment delivered (m3/km2/yr) = soil loss (t/ha/yr) x 100 / density x SDR; density 1.5 t/m3 (cell M2, Block 6 value).",
         "Catchments = union of the Basin 2 subbasins upstream of each dam junction (W3/sediment/catchments/Dam_catchments.shp). Hawasinah EX includes the 7 km2 of former S111 merged into S16 (engineer's decision, 2026-10-05)."]
for k, t in enumerate(notes): ws.cell(row=n + 1 + k, column=1, value=t).font = GREY
ws.freeze_panes = "B2"
out = SED / "B4_DAM_RUSLE.xlsx"; wb.save(out); print("written", out)
